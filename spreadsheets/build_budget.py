"""Builds Ultimate_Budget_Planner.xlsx — annual / monthly / period budget workbook."""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.workbook.defined_name import DefinedName

from common import *

D = dt.date
YEAR = 2026
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

REC0, RECN = 25, 60                 # Setup recurring rows 25..84
REC1 = REC0 + RECN - 1
PAY0, PAYN = 6, 1200                # Payments rows 6..1205
PAY1 = PAY0 + PAYN - 1
LOG0, LOGN = 6, 1000                # Manual Log rows 6..1005
LOG1 = LOG0 + LOGN - 1
GRID = "Calc!$B$2:$BB$61"           # occurrence grid (60 items x 53 occurrences)

P = lambda c: f"Payments!${c}${PAY0}:${c}${PAY1}"
L = lambda c: f"'Manual Log'!${c}${LOG0}:${c}${LOG1}"
S = lambda c: f"Setup!${c}${REC0}:${c}${REC1}"

START = "Setup!$D$5"
OPENING = "Setup!$D$6"
AUTOPAID = "Setup!$D$7"
YEAREND = "Setup!$D$9"

wb = Workbook()
wb.remove(wb.active)

sheets = {}
for name in (["Instructions", "Setup", "Payments", "Manual Log", "Accounts", "Dashboard"] + MONTHS
             + ["Annual", "Calendar", "Savings", "Debt Payoff", "Debt Schedule", "Lists", "Calc"]):
    sheets[name] = wb.create_sheet(name)

# ======================================================================== Lists
ws = sheets["Lists"]
lists = {
    "A": ["Category", "Income", "Bills", "Debt", "Subscriptions", "Savings"],
    "B": ["Log category", "Income", "Bills", "Debt", "Subscriptions", "Savings", "Variable Exp", "Transfer"],
    "C": ["Frequency", "Weekly", "Bi-weekly", "4-weekly", "Monthly", "Bi-monthly", "Quarterly",
          "Semi-annual", "Annual", "Once"],
    "D": ["Days", 7, 14, 28, 0, 0, 0, 0, 0, 0],
    "E": ["Months", 0, 0, 0, 1, 2, 3, 6, 12, 0],
    "F": ["Mark", "✓", "✗"],
    "I": ["Period", "Weekly", "Bi-weekly", "4-weekly", "Monthly", "Quarterly", "6-month", "Annual", "Custom"],
    "J": ["Debt method", "Debt Snowball", "Debt Avalanche", "Custom", "Credit Score Focus"],
    "K": ["YesNo", "Yes", "No"],
    "L": ["Week start", "Sunday", "Monday"],
}
for c, vals in lists.items():
    for i, v in enumerate(vals):
        ws[f"{c}{i+1}"] = v
ws["G1"] = "Accounts"
ws["G2"] = "N/A"
for i in range(12):
    ws[f"G{i+3}"] = f'=IF(Setup!F{6+i}="","",Setup!F{6+i})'
ws["H1"] = "Sub-categories"
for i in range(15):
    ws[f"H{i+2}"] = f'=IF(Setup!I{6+i}="","",Setup!I{6+i})'
for i in range(RECN):
    ws[f"H{i+17}"] = f'=IF(Setup!D{REC0+i}="","",Setup!D{REC0+i})'
ws.sheet_state = "hidden"
ACCTS = "=Lists!$G$2:$G$14"
SUBS = "=Lists!$H$2:$H$76"

# ======================================================================== Setup
ws = sheets["Setup"]
setup_sheet(ws, {"A": 2, "B": 5, "C": 18, "D": 24, "E": 12, "F": 16, "G": 14, "H": 13, "I": 20,
                 "J": 16, "K": 26, "L": 12, "M": 6, "N": 6, "O": 6, "P": 11}, tab=AMBER, rows=90)
merge_put(ws, 2, 2, 2, 11, "BUDGET SETUP", font(TEXT, 20, True), BG, align=LEFT)
put(ws, 3, 2, "Yellow cells are inputs. Fill this sheet once — Payments, Calendar, every monthly tab "
              "and the dashboards update automatically.", font(MUTED, 9, italic=True))
band(ws, 4, 3, 4, "GENERAL", LAVENDER)
for r, (lab, val, fmt) in enumerate([
        ("Budget start date", D(YEAR, 1, 1), DATE), ("Opening balance", 2500, MONEY),
        ("Auto-mark paid", "Yes", None), ("", None, None)], start=5):
    if lab:
        put(ws, r, 3, lab, font(MUTED), PANEL, align=LEFT)
        input_cell(ws, r, 4, val, fmt)
put(ws, 9, 3, "Year end", font(MUTED), PANEL, align=LEFT)
put(ws, 9, 4, f"=EDATE(D5,12)-1", font(TEXT, 10, True), PANEL, DATE, CENTER)
put(ws, 10, 3, "Auto-mark paid = Yes counts a scheduled item as paid once its date has passed "
               "(mark ✗ in Payments to skip one).", font(MUTED, 8, italic=True))
put(ws, 11, 3, "Opening balance = cash available at the start of month 1.", font(MUTED, 8, italic=True))
dv_list(ws, "=Lists!$K$2:$K$3", "D7")

band(ws, 4, 6, 7, "ACCOUNTS", TEAL)
put(ws, 5, 6, "Account", font(MUTED, 9, True), PANEL, align=LEFT)
put(ws, 5, 7, "Opening bal.", font(MUTED, 9, True), PANEL, align=CENTER)
accounts = [("Checking 1", 1500), ("Checking 2", 1000), ("Savings 1", 3000), ("Savings 2", 500),
            ("Credit Card 1", -450), ("Credit Card 2", -800), ("Credit Card 3", -300),
            ("Car Loan 1", -8000), ("Capital One", -2400)]
for i in range(12):
    a = accounts[i] if i < len(accounts) else (None, None)
    input_cell(ws, 6 + i, 6, a[0], align=LEFT)
    input_cell(ws, 6 + i, 7, a[1], MONEY)

band(ws, 4, 9, 10, "VARIABLE EXPENSES", LAVENDER)
put(ws, 5, 9, "Category", font(MUTED, 9, True), PANEL, align=LEFT)
put(ws, 5, 10, "Monthly budget", font(MUTED, 9, True), PANEL, align=CENTER)
variable = [("Groceries", 600), ("Eating Out", 400), ("Shopping", 300), ("Entertainment", 150),
            ("Transportation", 250), ("School", 200), ("Pets", 80), ("Hair/Beauty", 60),
            ("Gifts", 100), ("Other", 100)]
for i in range(15):
    v = variable[i] if i < len(variable) else (None, None)
    input_cell(ws, 6 + i, 9, v[0], align=LEFT)
    input_cell(ws, 6 + i, 10, v[1], MONEY)

band(ws, 23, 2, 11, "RECURRING TRANSACTIONS  —  set up once, auto-repeats into Payments, Calendar & all tabs",
     PINK)
heads = ["#", "Category", "Sub-category", "Amount", "Frequency", "Start date", "End date (opt.)",
         "From account", "To account", "Notes"]
for i, h in enumerate(heads):
    put(ws, 24, 2 + i, h, font(MUTED, 9, True), PANEL, align=CENTER, border=BOTTOM)
for i, h in enumerate(["key", "days", "mths", "n0", "end"]):
    put(ws, 24, 12 + i, h, font(MUTED, 8), BG)
rec = [
    ("Income", "Sam's Salary", 1600, "Bi-weekly", D(YEAR, 1, 2), None, "N/A", "Checking 1"),
    ("Income", "Lisa's Salary", 3200, "Monthly", D(YEAR, 1, 14), None, "N/A", "Checking 2"),
    ("Income", "Side Hustle", 450, "Monthly", D(YEAR, 1, 6), None, "N/A", "Checking 1"),
    ("Bills", "Rent", 1600, "Monthly", D(YEAR, 1, 1), None, "Checking 1", "N/A"),
    ("Bills", "Electricity Bill", 140, "Monthly", D(YEAR, 1, 2), None, "Credit Card 1", "N/A"),
    ("Bills", "Gas Bill", 50, "Monthly", D(YEAR, 1, 2), None, "Credit Card 1", "N/A"),
    ("Bills", "Water Bill", 60, "Monthly", D(YEAR, 1, 3), None, "Credit Card 3", "N/A"),
    ("Bills", "Internet", 75, "Monthly", D(YEAR, 1, 9), None, "Credit Card 2", "N/A"),
    ("Bills", "Phone Bill", 75, "Monthly", D(YEAR, 1, 9), None, "Credit Card 1", "N/A"),
    ("Bills", "Health Insurance", 150, "Monthly", D(YEAR, 1, 2), None, "Credit Card 2", "N/A"),
    ("Bills", "Car Insurance", 60, "Monthly", D(YEAR, 1, 6), None, "Credit Card 2", "N/A"),
    ("Bills", "Home Insurance", 55, "Monthly", D(YEAR, 1, 11), None, "Credit Card 2", "N/A"),
    ("Bills", "Childcare", 100, "Weekly", D(YEAR, 1, 7), None, "Checking 1", "N/A"),
    ("Bills", "Gym Membership", 75, "Monthly", D(YEAR, 1, 15), None, "Credit Card 1", "N/A"),
    ("Bills", "Car Registration", 320, "Annual", D(YEAR, 3, 10), None, "Checking 1", "N/A"),
    ("Debt", "Credit Card 1", 120, "Monthly", D(YEAR, 1, 10), None, "Checking 1", "Credit Card 1"),
    ("Debt", "Credit Card 2", 50, "Monthly", D(YEAR, 1, 8), None, "Checking 1", "Credit Card 2"),
    ("Debt", "Credit Card 3", 80, "Monthly", D(YEAR, 1, 4), None, "Checking 1", "Credit Card 3"),
    ("Debt", "Car Loan 1", 150, "Monthly", D(YEAR, 1, 12), None, "Checking 1", "Car Loan 1"),
    ("Debt", "Capital One", 100, "Monthly", D(YEAR, 1, 10), None, "Checking 1", "Capital One"),
    ("Subscriptions", "Netflix", 17, "Monthly", D(YEAR, 1, 12), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Spotify", 12, "Monthly", D(YEAR, 1, 15), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Amazon Prime", 15, "Monthly", D(YEAR, 1, 6), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Hulu", 12, "Monthly", D(YEAR, 1, 10), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Dropbox", 5, "Monthly", D(YEAR, 1, 22), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Canva", 25, "Monthly", D(YEAR, 1, 7), None, "Credit Card 2", "N/A"),
    ("Subscriptions", "Hello Fresh", 60, "Weekly", D(YEAR, 1, 6), None, "Credit Card 1", "N/A"),
    ("Savings", "Holiday", 200, "Monthly", D(YEAR, 1, 6), None, "Checking 1", "Savings 1"),
    ("Savings", "Emergency Fund", 300, "Monthly", D(YEAR, 1, 15), None, "Checking 2", "Savings 1"),
    ("Savings", "New Laptop", 150, "Monthly", D(YEAR, 1, 15), D(YEAR, 10, 31), "Checking 2", "Savings 2"),
    ("Savings", "Home Repairs", 150, "Monthly", D(YEAR, 1, 20), None, "Checking 2", "Savings 1"),
    ("Savings", "Wedding", 200, "Monthly", D(YEAR, 1, 6), None, "Checking 1", "Savings 2"),
]
for i in range(RECN):
    r = REC0 + i
    put(ws, r, 2, i + 1, font(MUTED, 8), PANEL, align=CENTER)
    v = rec[i] if i < len(rec) else (None,) * 8
    input_cell(ws, r, 3, v[0], align=LEFT)
    input_cell(ws, r, 4, v[1], align=LEFT)
    input_cell(ws, r, 5, v[2], MONEY)
    input_cell(ws, r, 6, v[3])
    input_cell(ws, r, 7, v[4], DATE)
    input_cell(ws, r, 8, v[5], DATE)
    input_cell(ws, r, 9, v[6], align=LEFT)
    input_cell(ws, r, 10, v[7], align=LEFT)
    input_cell(ws, r, 11, None, align=LEFT)
    ws.cell(r, 12).value = f'=IF(C{r}="","",C{r}&"|"&COUNTIF($C${REC0}:C{r},C{r}))'
    ws.cell(r, 13).value = f'=IFERROR(VLOOKUP(F{r},Lists!$C$2:$E$10,2,0),0)'
    ws.cell(r, 14).value = f'=IFERROR(VLOOKUP(F{r},Lists!$C$2:$E$10,3,0),0)'
    md = f"((YEAR($D$5)-YEAR(G{r}))*12+MONTH($D$5)-MONTH(G{r}))"
    a = f"MAX(0,ROUNDUP({md}/N{r},0))"
    ws.cell(r, 15).value = (f'=IF(OR(G{r}="",F{r}="",F{r}="Once"),0,IF(M{r}>0,MAX(0,ROUNDUP(($D$5-G{r})/M{r},0)),'
                            f'IF(N{r}=0,0,{a}+IF(EDATE(G{r},{a}*N{r})<$D$5,1,0))))')
    ws.cell(r, 16).value = f'=IF(H{r}="",$D$9,MIN(H{r},$D$9))'
    for c in range(12, 17):
        ws.cell(r, c).font = font(MUTED, 8)
        ws.cell(r, c).fill = fill(BG)
    ws.cell(r, 16).number_format = DATE
for c in "LMNOP":
    ws.column_dimensions[c].hidden = True
dv_list(ws, "=Lists!$A$2:$A$6", f"C{REC0}:C{REC1}")
dv_list(ws, "=Lists!$C$2:$C$10", f"F{REC0}:F{REC1}")
dv_list(ws, ACCTS, f"I{REC0}:J{REC1}")
for cat, colr in CAT_COLORS.items():
    cf_formula(ws, f"C{REC0}:C{REC1}", f'$C{REC0}="{cat}"', DARK, colr)
ws.freeze_panes = "A25"

# ======================================================================== Calc (occurrence grid)
ws = sheets["Calc"]
ws["A1"] = "item \\ k"
for k in range(53):
    ws.cell(1, 2 + k, k)
for i in range(RECN):
    r = 2 + i
    sr = REC0 + i
    ws.cell(r, 1, i + 1)
    for k in range(53):
        c = col(2 + k)
        date = (f"IF(Setup!$M${sr}>0,Setup!$G${sr}+(Setup!$O${sr}+{c}$1)*Setup!$M${sr},"
                f"EDATE(Setup!$G${sr},(Setup!$O${sr}+{c}$1)*Setup!$N${sr}))")
        ws.cell(r, 2 + k).value = (
            f'=IF(OR(Setup!$E${sr}="",Setup!$G${sr}="",Setup!$F${sr}=""),"",'
            f'IF(AND(Setup!$F${sr}="Once",{c}$1>0),"",'
            f'IF(OR({date}>Setup!$P${sr},{date}<{START}),"",{date}+$A{r}/1000)))')
ws.sheet_state = "hidden"

# ======================================================================== Payments
ws = sheets["Payments"]
setup_sheet(ws, {"A": 4, "B": 4, "C": 7, "D": 11, "E": 20, "F": 12, "G": 15, "H": 15, "I": 14,
                 "J": 2, "K": 11, "L": 12, "M": 15, "N": 15, "O": 2, "P": 11, "Q": 12, "R": 12,
                 "S": 12, "T": 11, "U": 12}, tab=PINK)
merge_put(ws, 2, 3, 2, 9, "PAYMENT SCHEDULE", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 3, "Auto-generated from Setup. Mark ✓ paid / ✗ skipped, use the filter arrows to sort & filter. "
              "Row colours follow the category.", font(MUTED, 9, italic=True))
band(ws, 4, 3, 9, "PAYMENT SCHEDULE", LAVENDER)
band(ws, 4, 11, 14, "VARIATION  —  enter any change for one payment here", AMBER)
band(ws, 4, 16, 20, "RESULT", TEAL)
for c, h in zip(range(1, 22), ["key", "idx", "PAID", "DATE", "SUB-CATEGORY", "AMOUNT", "FROM ACCOUNT",
                                "TO ACCOUNT", "CATEGORY", "", "NEW DATE", "NEW AMOUNT", "NEW FROM",
                                "NEW TO", "", "EFF. DATE", "ACTUAL PAID", "EFF. FROM", "EFF. TO", "STATUS",
                                "SCHEDULED"]):
    if h:
        put(ws, 5, c, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
for i in range(PAYN):
    r = PAY0 + i
    ws.cell(r, 1).value = f'=IFERROR(SMALL({GRID},{i+1}),"")'
    ws.cell(r, 2).value = f'=IF(A{r}="","",ROUND(MOD(A{r},1)*1000,0))'
    ws.cell(r, 4).value = f'=IF(A{r}="","",INT(A{r}))'
    ws.cell(r, 5).value = f'=IF(B{r}="","",INDEX({S("D")},B{r}))'
    ws.cell(r, 6).value = f'=IF(B{r}="","",INDEX({S("E")},B{r}))'
    ws.cell(r, 7).value = f'=IF(B{r}="","",INDEX({S("I")},B{r})&"")'
    ws.cell(r, 8).value = f'=IF(B{r}="","",INDEX({S("J")},B{r})&"")'
    ws.cell(r, 9).value = f'=IF(B{r}="","",INDEX({S("C")},B{r}))'
    ws.cell(r, 16).value = f'=IF(A{r}="","",IF(K{r}<>"",K{r},D{r}))'
    ws.cell(r, 17).value = (f'=IF(A{r}="",0,IF(C{r}="✗",0,IF(OR(C{r}="✓",AND({AUTOPAID}="Yes",P{r}<=TODAY())),'
                            f'IF(L{r}<>"",L{r},F{r}),0)))')
    ws.cell(r, 18).value = f'=IF(A{r}="","",IF(M{r}<>"",M{r},G{r}))'
    ws.cell(r, 19).value = f'=IF(A{r}="","",IF(N{r}<>"",N{r},H{r}))'
    ws.cell(r, 20).value = f'=IF(A{r}="","",IF(C{r}="✗","Skipped",IF(Q{r}>0,"Paid","Upcoming")))'
    ws.cell(r, 21).value = f'=IF(A{r}="",0,IF(C{r}="✗",0,IF(L{r}<>"",L{r},F{r})))'
    for c in range(1, 22):
        cell = ws.cell(r, c)
        cell.font = font(TEXT, 9)
        cell.fill = fill(ROW if i % 2 else ROW2)
        cell.alignment = CENTER
    for c in (4, 11, 16):
        ws.cell(r, c).number_format = DATE
    for c in (6, 12, 17, 21):
        ws.cell(r, c).number_format = MONEY
    for c in (3, 11, 12, 13, 14):
        ws.cell(r, c).fill = fill("2E2A1E")
    ws.cell(r, 5).alignment = LEFT
for c in "ABRSU":
    ws.column_dimensions[c].hidden = True
dv_list(ws, "=Lists!$F$2:$F$3", f"C{PAY0}:C{PAY1}", prompt="✓ = paid, ✗ = skipped")
dv_list(ws, ACCTS, f"M{PAY0}:N{PAY1}")
rng = f"C{PAY0}:I{PAY1}"
cf_formula(ws, f"C{PAY0}:U{PAY1}", f'$A{PAY0}=""', BG, BG, stop=True)
cf_formula(ws, f"T{PAY0}:T{PAY1}", f'$T{PAY0}="Paid"', TEAL)
cf_formula(ws, f"T{PAY0}:T{PAY1}", f'$T{PAY0}="Skipped"', MUTED)
for cat, colr in (("Income", TEAL), ("Savings", AMBER), ("Debt", CORAL)):
    cf_formula(ws, rng, f'$I{PAY0}="{cat}"', DARK, colr)
ws.auto_filter.ref = f"C5:T{PAY1}"
ws.freeze_panes = "D6"
wb.defined_names["PayCat"] = DefinedName("PayCat", attr_text=f"Payments!$I${PAY0}:$I${PAY1}")
wb.defined_names["PayAct"] = DefinedName("PayAct", attr_text=f"Payments!$Q${PAY0}:$Q${PAY1}")

# ======================================================================== Manual Log
ws = sheets["Manual Log"]
setup_sheet(ws, {"A": 2, "B": 12, "C": 15, "D": 20, "E": 12, "F": 15, "G": 15, "H": 36}, tab=LAVENDER)
merge_put(ws, 2, 2, 2, 8, "MANUAL LOG", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 2, "Log variable expenses, irregular income, one-off bills and money transfers here "
              "(or paste from your bank). Category 'Transfer' moves money between accounts only.",
    font(MUTED, 9, italic=True))
band(ws, 4, 2, 8, "MANUAL LOG", LAVENDER)
for i, h in enumerate(["DATE", "CATEGORY", "SUB-CATEGORY", "AMOUNT", "FROM ACCOUNT", "TO ACCOUNT",
                       "DESCRIPTION"]):
    put(ws, 5, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
random.seed(7)
log = []
spend = {"Groceries": (6, 45, 140), "Eating Out": (6, 18, 85), "Shopping": (3, 25, 140),
         "Entertainment": (2, 20, 70), "Transportation": (4, 30, 75), "School": (2, 30, 120),
         "Pets": (1, 25, 70), "Hair/Beauty": (1, 30, 60), "Gifts": (1, 20, 90), "Other": (2, 10, 50)}
cards = ["Credit Card 1", "Credit Card 2", "Checking 1", "Checking 2"]
for m in range(1, 10):
    for cat, (n, lo, hi) in spend.items():
        for _ in range(n):
            day = random.randint(1, 28)
            log.append((D(YEAR, m, day), "Variable Exp", cat, round(random.uniform(lo, hi), 2),
                        random.choice(cards), "N/A", ""))
    if m % 3 == 0:
        log.append((D(YEAR, m, 18), "Income", "Freelance Job", 1150, "N/A", "Checking 1", "Logo project"))
    log.append((D(YEAR, m, 25), "Transfer", "", 300, "Checking 2", "Checking 1", "Move to main account"))
log.sort()
for i in range(LOGN):
    r = LOG0 + i
    v = log[i] if i < len(log) else (None,) * 7
    fmts = [DATE, None, None, MONEY, None, None, None]
    for j in range(7):
        cell = put(ws, r, 2 + j, v[j], font(TEXT, 9), ROW if i % 2 else ROW2, fmts[j],
                   LEFT if j in (2, 6) else CENTER)
dv_list(ws, "=Lists!$B$2:$B$8", f"C{LOG0}:C{LOG1}")
dv_list(ws, SUBS, f"D{LOG0}:D{LOG1}")
dv_list(ws, ACCTS, f"F{LOG0}:G{LOG1}")
for cat, colr in (("Income", TEAL), ("Savings", AMBER), ("Debt", CORAL), ("Transfer", BLUE)):
    cf_formula(ws, f"B{LOG0}:H{LOG1}", f'$C{LOG0}="{cat}"', DARK, colr)
ws.auto_filter.ref = f"B5:H{LOG1}"
ws.freeze_panes = "B6"

# ======================================================================== Accounts
ws = sheets["Accounts"]
setup_sheet(ws, {"A": 2, "B": 20, "C": 14, "D": 14, "E": 14, "F": 16}, tab=TEAL)
merge_put(ws, 2, 2, 2, 6, "ACCOUNTS", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 2, "Balances update from paid items in Payments + everything in the Manual Log.",
    font(MUTED, 9, italic=True))
band(ws, 4, 2, 6, "ACCOUNT BALANCES", TEAL)
for i, h in enumerate(["ACCOUNT", "OPENING", "MONEY IN", "MONEY OUT", "CURRENT BALANCE"]):
    put(ws, 5, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
for i in range(12):
    r = 6 + i
    put(ws, r, 2, f'=IF(Setup!F{6+i}="","",Setup!F{6+i})', font(TEXT, 10, True), ROW, align=LEFT)
    put(ws, r, 3, f'=IF(B{r}="","",Setup!G{6+i})', font(), ROW, MONEY, CENTER)
    put(ws, r, 4, f'=IF(B{r}="","",SUMIFS({P("Q")},{P("S")},B{r})+SUMIFS({L("E")},{L("G")},B{r}))',
        font(TEAL), ROW, MONEY, CENTER)
    put(ws, r, 5, f'=IF(B{r}="","",SUMIFS({P("Q")},{P("R")},B{r})+SUMIFS({L("E")},{L("F")},B{r}))',
        font(PINK), ROW, MONEY, CENTER)
    put(ws, r, 6, f'=IF(B{r}="","",C{r}+D{r}-E{r})', font(TEXT, 11, True), ROW, MONEY, CENTER)
put(ws, 18, 2, "NET WORTH", font(DARK, 10, True), LAVENDER, align=LEFT)
for c in range(3, 7):
    put(ws, 18, c, f"=SUM({col(c)}6:{col(c)}17)", font(DARK, 10, True), LAVENDER, MONEY, CENTER)
cf_formula(ws, "F6:F17", "AND(ISNUMBER(F6),F6<0)", CORAL)
ch = BarChart(); ch.type = "bar"
ch.add_data(Reference(ws, min_col=6, min_row=5, max_row=17), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=6, max_row=17))
dark_chart(ch, legend=False); color_series(ch.series[0], TEAL)
ch.series[0].invertIfNegative = False
ch.width, ch.height = 14, 8
ws.add_chart(ch, "H4")


# ======================================================================== Dashboards
def act(cat, name, s="$C$6", e="$C$7"):
    return (f'SUMIFS({P("Q")},{P("I")},"{cat}",{P("E")},{name},{P("P")},">="&{s},{P("P")},"<="&{e})'
            f'+SUMIFS({L("E")},{L("C")},"{cat}",{L("D")},{name},{L("B")},">="&{s},{L("B")},"<="&{e})')


TABLES = {  # name: (category, first col, top row, items, idx helper col, kind)
    "INCOME": ("Income", 2, 27, 10, "AC", "in"),
    "SAVINGS": ("Savings", 7, 27, 10, "AD", "in"),
    "BILLS": ("Bills", 12, 27, 15, "AE", "out"),
    "DEBT": ("Debt", 17, 27, 10, "AF", "out"),
    "SUBSCRIPTIONS": ("Subscriptions", 2, 48, 10, "AG", "out"),
    "VARIABLE EXPENSES": ("Variable Exp", 7, 48, 15, None, "out"),
}


def build_dash(ws, kind, m=None, prev=None):
    setup_sheet(ws, {"A": 2, "B": 20, "C": 12, "D": 12, "E": 12, "F": 2, "G": 20, "H": 12, "I": 12,
                     "J": 12, "K": 2, "L": 20, "M": 12, "N": 12, "O": 12, "P": 2, "Q": 20, "R": 12,
                     "S": 12, "T": 12, "U": 2}, tab=TEAL if kind == "period" else LAVENDER,
                zoom=80, rows=68)
    ws.row_dimensions[3].height = 34
    # ---- header block
    paint(ws, 2, 2, 10, 5, PANEL)
    put(ws, 2, 2, "BUDGET DASHBOARD" if kind == "period" else "MONTHLY DASHBOARD",
        font(MUTED, 9, True), PANEL, align=LEFT)
    ws.merge_cells("B3:E3")
    lab = lambda r, t: put(ws, r, 2, t, font(MUTED, 9), PANEL, align=LEFT)
    if kind == "period":
        put(ws, 3, 2, '=UPPER(C5)&" DASHBOARD"', font(TEXT, 20, True), PANEL, align=LEFT)
        lab(4, "")
        put(ws, 4, 2, '=TEXT(C6,"d mmm yyyy")&"  –  "&TEXT(C7,"d mmm yyyy")', font(TEAL, 10, True),
            PANEL, align=LEFT)
        lab(5, "Budget period"); input_cell(ws, 5, 3, "Monthly")
        dv_list(ws, "=Lists!$I$2:$I$9", "C5")
        lab(6, "Start date"); input_cell(ws, 6, 3, D(YEAR, 9, 1), DATE)
        lab(7, "End date")
        put(ws, 7, 3, '=IF(C5="Custom",IF(C8="",C6,C8),IF(C5="Weekly",C6+6,IF(C5="Bi-weekly",C6+13,'
                      'IF(C5="4-weekly",C6+27,IF(C5="Monthly",EDATE(C6,1)-1,IF(C5="Quarterly",EDATE(C6,3)-1,'
                      'IF(C5="6-month",EDATE(C6,6)-1,EDATE(C6,12)-1)))))))', font(TEXT, 10, True), PANEL,
            DATE, CENTER)
        lab(8, "Custom end date"); input_cell(ws, 8, 3, None, DATE)
        put(ws, 8, 4, "← only for 'Custom'", font(MUTED, 8, italic=True), PANEL)
        before = lambda src, cat_rng, amt, date: (
            f'SUMIFS({amt},{cat_rng},"Income",{date},">="&{START},{date},"<"&C6)'
            f'-SUMIFS({amt},{cat_rng},"<>Income",{cat_rng},"<>Transfer",{date},">="&{START},{date},"<"&C6)')
        start_bal = (f'=IF(C9<>"",C9,{OPENING}+{before(0, P("I"), P("Q"), P("P"))}'
                     f'+{before(0, L("C"), L("E"), L("B"))})')
    else:
        put(ws, 3, 2, '=TEXT(C6,"mmmm yyyy")', font(TEXT, 22, True), PANEL, align=LEFT)
        put(ws, 4, 2, "Set the budget start date on the Setup sheet.", font(MUTED, 8, italic=True), PANEL,
            align=LEFT)
        lab(5, "Month #"); put(ws, 5, 3, m, font(MUTED), PANEL, align=CENTER)
        lab(6, "Start date")
        put(ws, 6, 3, f"=EDATE({START},C5-1)", font(TEXT, 10, True), PANEL, DATE, CENTER)
        lab(7, "End date")
        put(ws, 7, 3, "=EDATE(C6,1)-1", font(TEXT, 10, True), PANEL, DATE, CENTER)
        lab(8, "Days"); put(ws, 8, 3, "=C7-C6+1", font(TEXT), PANEL, align=CENTER)
        start_bal = (f'=IF(C9<>"",C9,{OPENING})' if prev is None else f"=IF(C9<>\"\",C9,'{prev}'!J8)")
    lab(9, "Start balance override"); input_cell(ws, 9, 3, None, MONEY)
    put(ws, 9, 4, "← leave blank to roll over", font(MUTED, 8, italic=True), PANEL)
    lab(10, "Start balance")
    put(ws, 10, 3, start_bal, font(TEAL, 11, True), PANEL, MONEY, CENTER)

    # ---- summary
    paint(ws, 2, 7, 10, 10, PANEL)
    band(ws, 2, 7, 10, "PERIOD SUMMARY" if kind == "period" else "MONTHLY SUMMARY", LAVENDER)
    put(ws, 3, 9, "BUDGET", font(MUTED, 8, True), PANEL, align=CENTER)
    put(ws, 3, 10, "ACTUAL", font(MUTED, 8, True), PANEL, align=CENTER)
    rows = [("Start balance", "=C10", "=C10", TEXT),
            ("+ Total income", "=C14", "=D14", TEAL),
            ("– Total expenses", "=SUM(C16:C19)", "=SUM(D16:D19)", PINK),
            ("– Total savings", "=C15", "=D15", AMBER),
            ("End balance", "=I4+I5-I6-I7", "=J4+J5-J6-J7", TEXT),
            ("Left to spend", "=I5-I6-I7", "=J5-J6-J7", LAVENDER)]
    for i, (t, b, a, c) in enumerate(rows):
        r = 4 + i
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        bold = t in ("End balance", "Left to spend")
        put(ws, r, 7, t, font(c, 10, bold), PANEL, align=LEFT)
        put(ws, r, 9, b, font(TEXT, 10, bold), PANEL, MONEY, CENTER)
        put(ws, r, 10, a, font(TEXT, 10, bold), PANEL, MONEY, CENTER)
    for c in (9, 10):
        ws.cell(8, c).border = Border(top=thin)

    # ---- KPI cards
    for (r, c, t, ref, colr) in [(2, 12, "INCOME", "J5", TEAL), (6, 12, "EXPENSES", "J6", PINK),
                                  (2, 17, "SAVINGS", "J7", AMBER), (6, 17, "LEFT TO SPEND", "J9", LAVENDER)]:
        band(ws, r, c, c + 3, t, colr, 10)
        merge_put(ws, r + 1, c, r + 2, c + 1, f"={ref}", font(colr, 20, True), PANEL, MONEY)
        put(ws, r + 1, c + 2, "BUDGET", font(MUTED, 8), PANEL, align=CENTER)
        put(ws, r + 1, c + 3, "PROGRESS", font(MUTED, 8), PANEL, align=CENTER)
        bref = ref.replace("J", "I")
        put(ws, r + 2, c + 2, f"={bref}", font(TEXT, 9), PANEL, MONEY, CENTER)
        put(ws, r + 2, c + 3, f"=IFERROR({ref}/{bref},0)", font(TEXT, 9, True), PANEL, PCT, CENTER)

    # ---- progress table
    band(ws, 12, 2, 5, "BUDGET vs ACTUAL", LAVENDER)
    paint(ws, 13, 2, 20, 5, PANEL)
    for i, h in enumerate(["CATEGORY", "BUDGET", "ACTUAL", "PROGRESS"]):
        put(ws, 13, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i else LEFT)
    prog = [("Income", "C30", "D30", TEAL), ("Savings", "H30", "I30", AMBER),
            ("Bills", "M30", "N30", PINK), ("Debt", "R30", "S30", CORAL),
            ("Subscriptions", "C51", "D51", ROSE), ("Variable Exp.", "H51", "I51", LAVENDER)]
    for i, (t, b, a, c) in enumerate(prog):
        r = 14 + i
        put(ws, r, 2, t, font(c, 10, True), PANEL, align=LEFT)
        put(ws, r, 3, f"={b}", font(), PANEL, MONEY, CENTER)
        put(ws, r, 4, f"={a}", font(), PANEL, MONEY, CENTER)
        put(ws, r, 5, f"=IFERROR(D{r}/C{r},0)", font(TEXT, 9, True), PANEL, PCT, CENTER)
        databar(ws, f"E{r}", c)
        if i >= 2:
            cf_formula(ws, f"D{r}", f"D{r}>C{r}", CORAL, bold=True)
        else:
            cf_formula(ws, f"D{r}", f"D{r}<C{r}", CORAL, bold=True)
    ws.merge_cells("B20:D20")
    put(ws, 20, 2, "REMAINING BUDGET (expenses)", font(MUTED, 9, True), PANEL, align=LEFT)
    put(ws, 20, 5, "=SUM(C16:C19)-SUM(D16:D19)", font(TEXT, 11, True), PANEL, MONEY, CENTER)
    cf_formula(ws, "E20", "E20<0", CORAL)
    put(ws, 22, 2, "Pink text = over budget (expenses) or below target (income & savings).",
        font(MUTED, 8, italic=True))
    put(ws, 23, 2, "Actual = paid items in Payments + Manual Log entries in this period.",
        font(MUTED, 8, italic=True))

    # ---- tables
    for title, (cat, c0, top, n, hcol, kd) in TABLES.items():
        colr = CAT_COLORS[cat]
        band(ws, top, c0, c0 + 3, title, colr)
        merge_put(ws, top + 1, c0, top + 1, c0 + 3, f"={col(c0+2)}{top+3}", font(colr, 16, True), PANEL, MONEY)
        ws.row_dimensions[top + 1].height = 24
        heads = ["SUB-CATEGORY", "BUDGET", "ACTUAL", "DIFF" if kd == "in" else "REMAINING"]
        for i, h in enumerate(heads):
            put(ws, top + 2, c0 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i else LEFT)
        first, last = top + 4, top + 3 + n
        nm, bu, ac, df = (col(c0 + i) for i in range(4))
        put(ws, top + 3, c0, "TOTAL", font(TEXT, 9, True), ROW, align=LEFT)
        for i, x in enumerate((bu, ac, df)):
            put(ws, top + 3, c0 + 1 + i, f"=SUM({x}{first}:{x}{last})", font(TEXT, 9, True), ROW, MONEY, CENTER)
        for k in range(1, n + 1):
            r = top + 3 + k
            if cat == "Variable Exp":
                sr = 5 + k
                name = f'=IF(Setup!$I${sr}="","",Setup!$I${sr})'
                if kind == "period":
                    budget = (f'=IF({nm}{r}="","",IF($C$5="Monthly",Setup!$J${sr},'
                              f'Setup!$J${sr}*($C$7-$C$6+1)*12/365))')
                else:
                    budget = f'=IF({nm}{r}="","",Setup!$J${sr})'
            else:
                ws[f"{hcol}{r}"] = f'=IFERROR(MATCH("{cat}|{k}",{S("L")},0),"")'
                name = f'=IF({hcol}{r}="","",INDEX({S("D")},{hcol}{r}))'
                budget = (f'=IF({hcol}{r}="","",COUNTIFS(INDEX({GRID},{hcol}{r},0),">="&$C$6,'
                          f'INDEX({GRID},{hcol}{r},0),"<"&($C$7+1))*INDEX({S("E")},{hcol}{r}))')
            actual = f'=IF({nm}{r}="","",{act(cat, f"{nm}{r}")})'
            diff = (f'=IF({nm}{r}="","",{ac}{r}-{bu}{r})' if kd == "in"
                    else f'=IF({nm}{r}="","",{bu}{r}-{ac}{r})')
            bg = PANEL if k % 2 else ROW2
            put(ws, r, c0, name, font(TEXT, 9), bg, align=LEFT)
            put(ws, r, c0 + 1, budget, font(TEXT, 9), bg, MONEY, CENTER)
            put(ws, r, c0 + 2, actual, font(TEXT, 9), bg, MONEY, CENTER)
            put(ws, r, c0 + 3, diff, font(TEXT, 9), bg, MONEY, CENTER)
        rng_a = f"{ac}{first}:{ac}{last}"
        if kd == "in":
            cf_formula(ws, rng_a, f'AND(ISNUMBER({ac}{first}),{ac}{first}<{bu}{first})', CORAL, bold=True)
        else:
            cf_formula(ws, rng_a, f'AND(ISNUMBER({ac}{first}),{ac}{first}>{bu}{first})', CORAL, bold=True)
        cf_formula(ws, f"{df}{top+3}:{df}{last}", f'AND(ISNUMBER({df}{top+3}),{df}{top+3}<0)', CORAL)

    # ---- top expenses (helper Y:AA rows 31..80)
    srcs = [("L", "N", 31, 15), ("Q", "S", 31, 10), ("B", "D", 52, 10), ("G", "I", 52, 15)]
    hr = 31
    for nc, vc, r0, n in srcs:
        for k in range(n):
            ws[f"Y{hr}"] = f"={nc}{r0+k}"
            ws[f"Z{hr}"] = f"={vc}{r0+k}"
            ws[f"AA{hr}"] = f'=IF(OR(Y{hr}="",N(Z{hr})<=0),"",Z{hr}+ROW()/1000000)'
            hr += 1
    band(ws, 48, 12, 15, "HIGHEST TO LOWEST EXPENSES", PINK)
    merge_put(ws, 49, 12, 49, 15, "=J6", font(PINK, 16, True), PANEL, MONEY)
    ws.row_dimensions[49].height = 24
    for i, h in enumerate(["CATEGORY", "AMOUNT", "% OF TOTAL", ""]):
        put(ws, 50, 12 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i else LEFT)
    for k in range(1, 11):
        r = 50 + k
        bg = PANEL if k % 2 else ROW2
        put(ws, r, 12, f'=IFERROR(INDEX($Y$31:$Y$80,MATCH(LARGE($AA$31:$AA$80,{k}),$AA$31:$AA$80,0)),"")',
            font(TEXT, 9), bg, align=LEFT)
        put(ws, r, 13, f'=IFERROR(INDEX($Z$31:$Z$80,MATCH(LARGE($AA$31:$AA$80,{k}),$AA$31:$AA$80,0)),"")',
            font(TEXT, 9), bg, MONEY, CENTER)
        put(ws, r, 14, f'=IFERROR(M{r}/$J$6,"")', font(TEXT, 9), bg, '0.0%', CENTER)
        put(ws, r, 15, None, font(), bg)
        databar(ws, f"N{r}", PINK)

    # ---- daily spending helper V:X rows 31..61
    ws["V29"] = "=MAX(1,ROUNDUP(($C$7-$C$6+1)/31,0))"
    for k in range(31):
        r = 31 + k
        ws[f"V{r}"] = f'=IF($C$6+{k}*$V$29>$C$7,"",$C$6+{k}*$V$29)'
        ws[f"W{r}"] = (f'=IF(V{r}="","",SUMIFS({P("Q")},{P("I")},"<>Income",{P("I")},"<>Savings",'
                       f'{P("P")},">="&V{r},{P("P")},"<"&(V{r}+$V$29),{P("P")},"<="&$C$7)'
                       f'+SUMIFS({L("E")},{L("C")},"<>Income",{L("C")},"<>Savings",{L("C")},"<>Transfer",'
                       f'{L("B")},">="&V{r},{L("B")},"<"&(V{r}+$V$29),{L("B")},"<="&$C$7))')
        ws[f"X{r}"] = f'=IF(V{r}="","",TEXT(V{r},"ddd d mmm"))'
    for c in ["V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG"]:
        ws.column_dimensions[c].hidden = True
    band(ws, 48, 17, 20, "DAILY SPENDING", LAVENDER)

    # ---- charts
    for (c, t, colr) in [(7, "BUDGET vs ACTUAL", LAVENDER), (12, "EXPENSES BREAKDOWN", PINK),
                         (17, "VARIABLE EXPENSES", LAVENDER)]:
        band(ws, 12, c, c + 3, t, colr)
    ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.overlap = -10; ch.gapWidth = 60
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=13, max_row=19), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=14, max_row=19))
    dark_chart(ch, legend_pos="t"); color_series(ch.series[0], "5C5C78"); color_series(ch.series[1], TEAL)
    ch.y_axis.numFmt = '$#,##0'; ch.x_axis.scaling.orientation = "maxMin"
    ch.width, ch.height = 10.6, 7.0
    ws.add_chart(ch, "G13")

    dn = DoughnutChart(holeSize=55)
    dn.add_data(Reference(ws, min_col=4, min_row=16, max_row=19))
    dn.set_categories(Reference(ws, min_col=2, min_row=16, max_row=19))
    dark_chart(dn, axes=False); color_points(dn.series[0], [PINK, CORAL, ROSE, LAVENDER]); pct_labels(dn)
    dn.width, dn.height = 10.6, 7.0
    ws.add_chart(dn, "L13")

    dn2 = DoughnutChart(holeSize=55)
    dn2.add_data(Reference(ws, min_col=9, min_row=52, max_row=66))
    dn2.set_categories(Reference(ws, min_col=7, min_row=52, max_row=66))
    dark_chart(dn2, axes=False); color_points(dn2.series[0], PALETTE); pct_labels(dn2)
    dn2.width, dn2.height = 10.6, 7.0
    ws.add_chart(dn2, "Q13")

    bc = BarChart(); bc.type = "bar"; bc.gapWidth = 40
    bc.add_data(Reference(ws, min_col=23, min_row=31, max_row=61))
    bc.set_categories(Reference(ws, min_col=24, min_row=31, max_row=61))
    dark_chart(bc, legend=False); color_series(bc.series[0], LAVENDER)
    bc.visible_cells_only = False
    bc.x_axis.scaling.orientation = "maxMin"; bc.y_axis.numFmt = '$#,##0'
    bc.width, bc.height = 10.6, 9.6
    ws.add_chart(bc, "Q49")
    ws.freeze_panes = "A2"


build_dash(sheets["Dashboard"], "period")
for i, mname in enumerate(MONTHS):
    build_dash(sheets[mname], "month", i + 1, MONTHS[i - 1] if i else None)

# ======================================================================== Annual
ws = sheets["Annual"]
setup_sheet(ws, dict([("A", 2), ("B", 22)] + [(col(3 + i), 11) for i in range(12)] +
                     [("O", 13), ("P", 12), ("Q", 13), ("R", 11)]), tab=AMBER, zoom=80, rows=150)
merge_put(ws, 2, 2, 2, 10, "ANNUAL BREAKDOWN", font(TEXT, 20, True), BG, align=LEFT)
put(ws, 3, 2, '=TEXT(Jan!C6,"mmm yyyy")&" – "&TEXT(Dec!C7,"mmm yyyy")&"   ·   actual totals pulled from the 12 monthly tabs"',
    font(MUTED, 9, italic=True))


def month_header(r, title, colr):
    band(ws, r, 2, 18, title, colr)
    put(ws, r + 1, 2, "CATEGORY", font(MUTED, 8, True), PANEL, align=LEFT)
    for i, mn in enumerate(MONTHS):
        put(ws, r + 1, 3 + i, f'=TEXT({mn}!$C$6,"mmm")', font(MUTED, 8, True), PANEL, align=CENTER)
    for i, h in enumerate(["TOTAL", "AVERAGE", "BUDGET", "% USED"]):
        put(ws, r + 1, 15 + i, h, font(MUTED, 8, True), PANEL, align=CENTER)


month_header(5, "ALL CATEGORIES — ACTUAL", LAVENDER)
cats = [("Income", 14, TEAL), ("Savings", 15, AMBER), ("Bills", 16, PINK), ("Debt", 17, CORAL),
        ("Subscriptions", 18, ROSE), ("Variable Exp.", 19, LAVENDER)]
for i, (t, pr, colr) in enumerate(cats):
    r = 7 + i
    put(ws, r, 2, t, font(colr, 10, True), PANEL, align=LEFT)
    for j, mn in enumerate(MONTHS):
        put(ws, r, 3 + j, f"={mn}!D{pr}", font(TEXT, 9), PANEL, MONEY0, CENTER)
    put(ws, r, 15, f"=SUM(C{r}:N{r})", font(TEXT, 10, True), ROW, MONEY0, CENTER)
    put(ws, r, 16, f'=IFERROR(O{r}/MAX(1,COUNTIF(C{r}:N{r},">0")),0)', font(TEXT, 9), ROW, MONEY0, CENTER)
    put(ws, r, 17, "=" + "+".join(f"{mn}!C{pr}" for mn in MONTHS), font(TEXT, 9), ROW, MONEY0, CENTER)
    put(ws, r, 18, f"=IFERROR(O{r}/Q{r},0)", font(TEXT, 9, True), ROW, PCT, CENTER)
for r, t, f, colr in [(13, "TOTAL EXPENSES", "SUM({c}9:{c}12)", PINK),
                      (14, "NET (income – exp. – savings)", "{c}7-{c}8-{c}13", TEAL)]:
    put(ws, r, 2, t, font(DARK, 9, True), colr, align=LEFT)
    for j in range(3, 19):
        c = col(j)
        fmt = PCT if j == 18 else MONEY0
        v = f"=IFERROR(O{r}/Q{r},0)" if j == 18 else "=" + f.format(c=c)
        if r == 14 and j == 18:
            v = None
        put(ws, r, j, v, font(DARK, 9, True), colr, fmt, CENTER)
# charts
band(ws, 16, 2, 10, "MONTHLY EXPENSES BY CATEGORY", PINK)
band(ws, 16, 12, 18, "WHERE THE MONEY WENT", LAVENDER)
sc = BarChart(); sc.type = "col"; sc.grouping = "stacked"; sc.overlap = 100; sc.gapWidth = 50
for i, r in enumerate(range(9, 13)):
    sc.add_data(Reference(ws, min_col=2, max_col=14, min_row=r), from_rows=True, titles_from_data=True)
sc.set_categories(Reference(ws, min_col=3, max_col=14, min_row=6))
dark_chart(sc, legend_pos="t", gridlines=True)
for s, c in zip(sc.series, [PINK, CORAL, ROSE, LAVENDER]):
    color_series(s, c)
lc = LineChart()
lc.add_data(Reference(ws, min_col=2, max_col=14, min_row=7), from_rows=True, titles_from_data=True)
color_series(lc.series[0], TEAL, line=True)
sc += lc
sc.y_axis.numFmt = '$#,##0'
sc.width, sc.height = 22, 8
ws.add_chart(sc, "B18")
dn = DoughnutChart(holeSize=55)
dn.add_data(Reference(ws, min_col=15, min_row=8, max_row=12))
dn.set_categories(Reference(ws, min_col=2, min_row=8, max_row=12))
dark_chart(dn, axes=False); color_points(dn.series[0], [AMBER, PINK, CORAL, ROSE, LAVENDER]); pct_labels(dn)
dn.width, dn.height = 13, 8
ws.add_chart(dn, "L18")
# breakdowns
r = 34
for title, (cat, c0, top, n, hcol, kd) in TABLES.items():
    colr = CAT_COLORS[cat]
    month_header(r, f"{title} — BY MONTH", colr)
    tr = r + 2
    put(ws, tr, 2, "TOTAL", font(TEXT, 9, True), ROW, align=LEFT)
    for k in range(n):
        rr = tr + 1 + k
        src = top + 4 + k
        put(ws, rr, 2, f"=Jan!{col(c0)}{src}", font(TEXT, 9), PANEL, align=LEFT)
        for j, mn in enumerate(MONTHS):
            put(ws, rr, 3 + j, f"={mn}!{col(c0+2)}{src}", font(TEXT, 9), PANEL, MONEY0, CENTER)
        put(ws, rr, 15, f'=IF($B{rr}="","",SUM(C{rr}:N{rr}))', font(TEXT, 9, True), ROW, MONEY0, CENTER)
        put(ws, rr, 16, f'=IF($B{rr}="","",IFERROR(O{rr}/MAX(1,COUNTIF(C{rr}:N{rr},">0")),0))', font(TEXT, 9),
            ROW, MONEY0, CENTER)
        put(ws, rr, 17, f'=IF($B{rr}="","",' + "+".join(f"N({mn}!{col(c0+1)}{src})" for mn in MONTHS) + ")",
            font(TEXT, 9), ROW, MONEY0, CENTER)
        put(ws, rr, 18, f'=IF($B{rr}="","",IFERROR(O{rr}/Q{rr},0))', font(TEXT, 9), ROW, PCT, CENTER)
    for j in range(3, 19):
        c = col(j)
        v = f"=IFERROR(O{tr}/Q{tr},0)" if j == 18 else f"=SUM({c}{tr+1}:{c}{tr+n})"
        put(ws, tr, j, v, font(TEXT, 9, True), ROW, PCT if j == 18 else MONEY0, CENTER)
    r = tr + n + 2
ws.freeze_panes = "C7"

# ======================================================================== Calendar
ws = sheets["Calendar"]
widths = {"A": 2}
for d in range(7):
    widths[col(2 + 3 * d)] = 15
    widths[col(3 + 3 * d)] = 9
    widths[col(4 + 3 * d)] = 3
widths.update({"W": 2, "X": 14, "Y": 12})
setup_sheet(ws, widths, tab=PINK, zoom=85, rows=45)
merge_put(ws, 2, 2, 3, 8, '=TEXT(R2,"mmmm yyyy")', font(TEXT, 24, True), BG, align=LEFT)
ws.merge_cells("B4:J4")
put(ws, 4, 2, "Items come from Payments · paid = grey · small number by each date = projected end-of-day balance",
    font(MUTED, 8, italic=True))
for r, (t, v, fmt) in enumerate([("Any date in month", D(YEAR, 9, 1), DATE), ("Week starts on", "Sunday", None),
                                 ("Start balance (optional)", None, MONEY)], start=2):
    ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
    put(ws, r, 11, t, font(MUTED, 9), PANEL, align=LEFT)
    input_cell(ws, r, 14, v, fmt)
dv_list(ws, "=Lists!$L$2:$L$3", "N3")
for r, (t, v, fmt) in enumerate([("Month start", "=DATE(YEAR(N2),MONTH(N2),1)", DATE),
                                 ("Grid start", '=R2-MOD(WEEKDAY(R2,IF(N3="Monday",2,1))-1,7)', DATE),
                                 ("Start balance", f'=IF(N4<>"",N4,{OPENING}+SUMIFS({P("U")},{P("I")},"Income",'
                                  f'{P("P")},">="&{START},{P("P")},"<"&R3)-SUMIFS({P("U")},{P("I")},"<>Income",'
                                  f'{P("P")},">="&{START},{P("P")},"<"&R3))', MONEY)], start=2):
    ws.merge_cells(start_row=r, start_column=16, end_row=r, end_column=17)
    put(ws, r, 16, t, font(MUTED, 9), PANEL, align=LEFT)
    put(ws, r, 18, v, font(TEXT, 9, True), PANEL, fmt, CENTER)
legend = [("INCOME", TEAL, 2, 24), ("BILLS & SUBS", LAVENDER, 2, 25), ("DEBT", CORAL, 3, 24),
          ("SAVINGS", AMBER, 3, 25), ("PAID", "4A4A5C", 4, 24)]
for t, c, r, cc in legend:
    put(ws, r, cc, t, font(DARK if c != "4A4A5C" else TEXT, 8, True), c, align=CENTER)
for d in range(7):
    c = 2 + 3 * d
    ws.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c + 1)
    put(ws, 6, c, f'=UPPER(TEXT($R$3+{d},"dddd"))', font(DARK, 10, True), LAVENDER, align=CENTER)
    ws.cell(6, c + 1).fill = fill(LAVENDER)
band(ws, 6, 24, 25, "WEEKLY BALANCE", LAVENDER)
for w in range(6):
    base = 7 + 6 * w
    ws.row_dimensions[base].height = 18
    for d in range(7):
        c = 2 + 3 * d
        h = col(c + 2)
        ws.cell(base, c + 2).value = f"=$R$3+{7*w+d}"
        ws.cell(base, c + 2).number_format = "0"
        put(ws, base, c, f"=DAY({h}{base})", font(TEXT, 11, True), "2A2A3A", align=LEFT)
        put(ws, base, c + 1, (f'=IF(MONTH({h}{base})<>MONTH($R$2),"",$R$4+SUMIFS({P("U")},{P("I")},"Income",'
                              f'{P("P")},">="&$R$3,{P("P")},"<="&{h}{base})-SUMIFS({P("U")},{P("I")},"<>Income",'
                              f'{P("P")},">="&$R$3,{P("P")},"<="&{h}{base}))'),
            font(MUTED, 8), "2A2A3A", MONEY0, RIGHT)
        cf_formula(ws, f"{col(c)}{base}", f"MONTH({h}{base})<>MONTH($R$2)", "5A5A70")
        for k in range(1, 6):
            r = base + k
            ws.cell(r, c + 2).value = (f'=IFERROR(_xlfn.AGGREGATE(15,6,(ROW({P("P")})-{PAY0-1})/'
                                       f'({P("P")}={h}${base}),{k}),"")')
            put(ws, r, c, f'=IF({h}{r}="","",INDEX({P("E")},{h}{r}))', font(TEXT, 8), PANEL, align=LEFT)
            put(ws, r, c + 1, f'=IF({h}{r}="","",INDEX({P("U")},{h}{r}))', font(TEXT, 8), PANEL, MONEY0, RIGHT)
        ws.cell(base + 5, c).border = Border(bottom=thin)
        ws.cell(base + 5, c + 1).border = Border(bottom=thin)
    ws0 = f"$D${base}"
    we = f"$V${base}"
    put(ws, base, 24, f'="WEEK {w+1}"', font(DARK, 9, True), LAVENDER, align=LEFT)
    put(ws, base, 25, f'=TEXT({ws0},"d mmm")&" – "&TEXT({we},"d mmm")', font(DARK, 8, True), LAVENDER,
        align=CENTER)
    sums = lambda crit: (f'SUMIFS({P("U")},{P("I")},{crit},{P("P")},">="&{ws0},{P("P")},"<="&{we})')
    items = [("+ Income", "=" + sums('"Income"'), TEAL),
             ("– Bills & Subs", f'={sums(chr(34)+"Bills"+chr(34))}+{sums(chr(34)+"Subscriptions"+chr(34))}', LAVENDER),
             ("– Debt", "=" + sums('"Debt"'), CORAL),
             ("– Savings", "=" + sums('"Savings"'), AMBER),
             ("End balance", f"=$R$4+SUMIFS({P('U')},{P('I')},\"Income\",{P('P')},\">=\"&$R$3,{P('P')},\"<=\"&{we})"
                             f"-SUMIFS({P('U')},{P('I')},\"<>Income\",{P('P')},\">=\"&$R$3,{P('P')},\"<=\"&{we})", TEXT)]
    for k, (t, f, c) in enumerate(items, start=1):
        put(ws, base + k, 24, t, font(c, 9, k == 5), PANEL, align=LEFT)
        put(ws, base + k, 25, f, font(TEXT, 9, k == 5), PANEL, MONEY0, CENTER)
for d in range(7):
    c = 2 + 3 * d
    ws.column_dimensions[col(c + 2)].hidden = True
    for (cc, hc) in ((c, c + 2), (c + 1, c + 2)):
        rng = f"{col(cc)}8:{col(cc)}42"
        h = f"${col(hc)}8"
        guard = f"AND(ISNUMBER({h}),{h}<5000"
        cf_formula(ws, rng, f"{guard},INDEX(PayAct,{h})>0)", "B8B8C8", "4A4A5C", stop=True)
        for cat, colr in (("Income", TEAL), ("Debt", CORAL), ("Savings", AMBER), ("Bills", LAVENDER),
                          ("Subscriptions", LAVENDER)):
            cf_formula(ws, rng, f'{guard},INDEX(PayCat,{h})="{cat}")', DARK, colr)
ws.freeze_panes = "A7"

# ======================================================================== Savings
ws = sheets["Savings"]
setup_sheet(ws, {"A": 2, "B": 20, "C": 12, "D": 12, "E": 12, "F": 13, "G": 13, "H": 12, "I": 11, "J": 10,
                 "K": 13, "L": 12}, tab=AMBER, rows=30)
merge_put(ws, 2, 2, 2, 8, "SINKING FUNDS / SAVINGS TRACKER", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 2, "Name each fund exactly like its Savings sub-category in Setup — contributions are pulled "
              "automatically from Payments & the Manual Log.", font(MUTED, 9, italic=True))
band(ws, 4, 2, 12, "SAVINGS GOALS", AMBER)
heads = ["FUND", "GOAL", "TARGET DATE", "ALREADY SAVED", "CONTRIBUTIONS", "TOTAL SAVED", "LEFT TO SAVE",
         "PROGRESS", "MONTHS LEFT", "NEEDED / MONTH", "STATUS"]
for i, h in enumerate(heads):
    put(ws, 5, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
funds = [("Holiday", 2400, D(YEAR, 12, 1), 300), ("Emergency Fund", 10000, D(YEAR + 1, 6, 30), 2000),
         ("New Laptop", 1800, D(YEAR, 10, 31), 0), ("Home Repairs", 1500, D(YEAR + 1, 3, 31), 200),
         ("Wedding", 12000, D(YEAR + 1, 9, 30), 3000)]
for i in range(20):
    r = 6 + i
    f = funds[i] if i < len(funds) else (None,) * 4
    input_cell(ws, r, 2, f[0], align=LEFT)
    input_cell(ws, r, 3, f[1], MONEY0)
    input_cell(ws, r, 4, f[2], DATE)
    input_cell(ws, r, 5, f[3], MONEY0)
    put(ws, r, 6, f'=IF(B{r}="","",SUMIFS({P("Q")},{P("I")},"Savings",{P("E")},B{r})'
                  f'+SUMIFS({L("E")},{L("C")},"Savings",{L("D")},B{r}))', font(), PANEL, MONEY0, CENTER)
    put(ws, r, 7, f'=IF(B{r}="","",N(E{r})+F{r})', font(AMBER, 10, True), PANEL, MONEY0, CENTER)
    put(ws, r, 8, f'=IF(B{r}="","",MAX(0,N(C{r})-G{r}))', font(), PANEL, MONEY0, CENTER)
    put(ws, r, 9, f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,G{r}/C{r}))', font(TEXT, 9, True), PANEL, PCT, CENTER)
    put(ws, r, 10, f'=IF(OR(B{r}="",D{r}=""),"",MAX(0,DATEDIF(MIN(TODAY(),D{r}),D{r},"m")))', font(), PANEL,
        "0", CENTER)
    put(ws, r, 11, f'=IF(OR(B{r}="",J{r}=""),"",IF(H{r}=0,0,H{r}/MAX(1,J{r})))', font(), PANEL, MONEY0, CENTER)
    put(ws, r, 12, f'=IF(B{r}="","",IF(G{r}>=N(C{r}),"✓ Reached",IF(AND(D{r}<>"",D{r}<TODAY()),"Overdue","Ongoing")))',
        font(TEXT, 9), PANEL, align=CENTER)
    databar(ws, f"I{r}", AMBER)
cf_formula(ws, "L6:L25", 'L6="✓ Reached"', TEAL, bold=True)
cf_formula(ws, "L6:L25", 'L6="Overdue"', CORAL, bold=True)
put(ws, 27, 2, "TOTAL", font(DARK, 10, True), AMBER, align=LEFT)
for c in (3, 5, 6, 7, 8):
    put(ws, 27, c, f"=SUM({col(c)}6:{col(c)}25)", font(DARK, 10, True), AMBER, MONEY0, CENTER)
put(ws, 27, 9, "=IFERROR(G27/C27,0)", font(DARK, 10, True), AMBER, PCT, CENTER)
for c in (4, 10, 11, 12):
    ws.cell(27, c).fill = fill(AMBER)
ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.overlap = 100; ch.gapWidth = 50
ch.add_data(Reference(ws, min_col=3, min_row=5, max_row=10), titles_from_data=True)
ch.add_data(Reference(ws, min_col=7, min_row=5, max_row=10), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=6, max_row=10))
dark_chart(ch, legend_pos="t"); color_series(ch.series[0], "3A3A4C"); color_series(ch.series[1], AMBER)
ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.numFmt = '$#,##0'
ch.width, ch.height = 16, 8
ws.add_chart(ch, "N4")

# ======================================================================== Debt Payoff
ws = sheets["Debt Payoff"]
setup_sheet(ws, {"A": 2, "B": 20, "C": 13, "D": 10, "E": 12, "F": 12, "G": 10, "H": 8, "I": 11, "J": 11,
                 "K": 12, "L": 12, "M": 11, "N": 4}, tab=CORAL, rows=45)
merge_put(ws, 2, 2, 2, 9, "DEBT PAYOFF PLANNER", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 2, "Pick a method, enter your debts and any extra monthly payment. Freed-up minimums roll into "
              "the next debt automatically (snowball effect).", font(MUTED, 9, italic=True))
band(ws, 4, 2, 3, "SETTINGS", CORAL)
for r, (t, v, fmt) in enumerate([("Method", "Debt Snowball", None), ("Start month", D(YEAR, 1, 1), DATE),
                                 ("Extra monthly payment", 200, MONEY0)], start=5):
    put(ws, r, 2, t, font(MUTED, 9), PANEL, align=LEFT)
    input_cell(ws, r, 3, v, fmt)
dv_list(ws, "=Lists!$J$2:$J$5", "C5")
band(ws, 4, 5, 9, "RESULTS", LAVENDER)
res = [("Debt-free date", "=IF(MAX(J15:J24)>300,\"300+ months\",EDATE(C6,MAX(J15:J24)-1))", 'mmmm yyyy'),
       ("Months to debt-free", "=MAX(J15:J24)", "0"),
       ("Total interest", "=SUM(L15:L24)", MONEY0),
       ("Total starting debt", "=SUM(C15:C24)", MONEY0),
       ("Monthly payment budget", "=C12", MONEY0)]
for i, (t, f, fmt) in enumerate(res):
    r = 5 + i
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
    put(ws, r, 5, t, font(MUTED, 9), PANEL, align=LEFT)
    put(ws, r, 7, f, font(TEXT if i else CORAL, 12 if i == 0 else 10, True), PANEL, fmt, CENTER)
put(ws, 12, 2, "Monthly budget (min. + extra)", font(MUTED, 9), PANEL, align=LEFT)
put(ws, 12, 3, "=SUM(E15:E24)+N(C7)", font(TEXT, 10, True), PANEL, MONEY0, CENTER)
band(ws, 13, 2, 13, "YOUR DEBTS  (up to 10)", CORAL)
heads = ["DEBT", "BALANCE", "APR %", "MIN. PAYMENT", "CREDIT LIMIT", "CUSTOM ORDER", "ORDER", "UTILISATION",
         "PAYOFF MONTH", "PAYOFF DATE", "TOTAL INTEREST", "PAID OFF %"]
for i, h in enumerate(heads):
    put(ws, 14, 2 + i, h, font(MUTED, 7, True), PANEL, align=CENTER, border=BOTTOM)
debts = [("Credit Card 1", 3200, 0.229, 120, 5000, 2), ("Credit Card 2", 1500, 0.199, 50, 3000, 1),
         ("Capital One", 2400, 0.249, 100, 4000, 3), ("Car Loan 1", 8000, 0.065, 150, None, 4),
         ("Student Loan", 12000, 0.045, 60, None, 5)]
for i in range(10):
    r = 15 + i
    dd = debts[i] if i < len(debts) else (None,) * 6
    input_cell(ws, r, 2, dd[0], align=LEFT)
    input_cell(ws, r, 3, dd[1], MONEY0)
    input_cell(ws, r, 4, dd[2], "0.0%")
    input_cell(ws, r, 5, dd[3], MONEY0)
    input_cell(ws, r, 6, dd[4], MONEY0)
    input_cell(ws, r, 7, dd[5], "0")
    ws[f"O{r}"] = (f'=IF(N(C{r})<=0,1E+9,CHOOSE(MATCH($C$5,Lists!$J$2:$J$5,0),C{r},-D{r},N(G{r}),'
                   f'-IFERROR(C{r}/F{r},0)))')
    put(ws, r, 8, f'=IF(N(C{r})<=0,"",COUNTIF($O$15:$O$24,"<"&O{r})+COUNTIF($O$15:O{r},O{r}))',
        font(CORAL, 10, True), PANEL, "0", CENTER)
    put(ws, r, 9, f'=IF(OR(N(C{r})<=0,N(F{r})=0),"",C{r}/F{r})', font(), PANEL, PCT, CENTER)
    bcol = col(3 + i)
    put(ws, r, 10, f"=IF(N(C{r})<=0,\"\",COUNTIF('Debt Schedule'!{bcol}$6:{bcol}$305,\">0.004\")+1)",
        font(), PANEL, "0", CENTER)
    put(ws, r, 11, f'=IF(J{r}="","",IF(J{r}>300,"300+",EDATE($C$6,J{r}-1)))', font(TEXT, 10, True), PANEL,
        'mmm yyyy', CENTER)
    icol = col(13 + i)
    put(ws, r, 12, f"=IF(N(C{r})<=0,\"\",SUM('Debt Schedule'!{icol}$6:{icol}$305))", font(), PANEL, MONEY0,
        CENTER)
    put(ws, r, 13, (f"=IF(N(C{r})<=0,\"\",1-INDEX('Debt Schedule'!{bcol}$5:{bcol}$305,"
                    f"MIN(301,MAX(1,DATEDIF($C$6,MAX($C$6,TODAY()),\"m\")+1)))/C{r})"), font(TEXT, 9, True),
        PANEL, PCT, CENTER)
    databar(ws, f"M{r}", CORAL)
ws.column_dimensions["O"].hidden = True
put(ws, 26, 2, "Snowball = smallest balance first · Avalanche = highest APR first · Custom = your order · "
               "Credit Score Focus = highest utilisation first.", font(MUTED, 8, italic=True))
ch = LineChart()
ch.add_data(Reference(sheets["Debt Schedule"], min_col=53, min_row=4, max_row=125), titles_from_data=True)
ch.set_categories(Reference(sheets["Debt Schedule"], min_col=2, min_row=5, max_row=125))
dark_chart(ch, legend=False, gridlines=True); color_series(ch.series[0], CORAL, line=True)
ch.y_axis.numFmt = '$#,##0'; ch.x_axis.tickLblSkip = 12
ch.width, ch.height = 24, 7.5
band(ws, 28, 2, 13, "TOTAL DEBT OVER TIME (first 10 years)", CORAL)
ws.add_chart(ch, "B30")

# ---- Debt Schedule
ws = sheets["Debt Schedule"]
setup_sheet(ws, dict([("A", 7), ("B", 10)] + [(col(c), 10) for c in range(3, 56)]), tab=CORAL, zoom=80)
put(ws, 1, 1, "DEBT SCHEDULE — calculated month by month from the Debt Payoff sheet (no inputs here).",
    font(TEXT, 11, True))
groups = [("BALANCE", CORAL), ("INTEREST", AMBER), ("MIN. PAID", LAVENDER), ("REMAINING NEED", "5C5C78"),
          ("TOTAL PAYMENT", TEAL)]
for g, (t, c) in enumerate(groups):
    band(ws, 2, 3 + 10 * g, 12 + 10 * g, t, c, 9)
    for j in range(10):
        cc = 3 + 10 * g + j
        put(ws, 3, cc, f"=IF('Debt Payoff'!$B${15+j}=\"\",\"—\",'Debt Payoff'!$B${15+j})", font(MUTED, 8, True),
            PANEL, align=CENTER)
        put(ws, 4, cc, f"=N('Debt Payoff'!$H${15+j})", font(MUTED, 8), PANEL, "0", CENTER)
for i, t in enumerate(["TOTAL BALANCE", "TOTAL PAID", "TOTAL INTEREST"]):
    put(ws, 3, 53 + i, t, font(MUTED, 8, True), PANEL, align=CENTER)
    put(ws, 4, 53 + i, t.title(), font(MUTED, 8), PANEL, align=CENTER)
put(ws, 4, 1, "Month", font(MUTED, 8, True), PANEL); put(ws, 4, 2, "Date / order →", font(MUTED, 8, True), PANEL)
put(ws, 5, 1, 0, font(MUTED, 8)); put(ws, 5, 2, "Start", font(MUTED, 8))
for j in range(10):
    ws.cell(5, 3 + j).value = f"=N('Debt Payoff'!$C${15+j})"
    ws.cell(5, 3 + j).number_format = MONEY0
ws.cell(5, 53).value = "=SUM(C5:L5)"
ws.cell(5, 53).number_format = MONEY0
for t in range(1, 301):
    r = 5 + t
    ws.cell(r, 1, t).font = font(MUTED, 8)
    ws.cell(r, 2).value = f"=EDATE('Debt Payoff'!$C$6,{t-1})"
    ws.cell(r, 2).number_format = "mmm yy"
    for j in range(10):
        b, i_, m_, n_, p_ = (col(3 + 10 * g + j) for g in range(5))
        apr = f"N('Debt Payoff'!$D${15+j})"
        mn = f"N('Debt Payoff'!$E${15+j})"
        ws[f"{i_}{r}"] = f"={b}{r-1}*{apr}/12"
        ws[f"{m_}{r}"] = f"=MIN({mn},{b}{r-1}+{i_}{r})"
        ws[f"{n_}{r}"] = f"={b}{r-1}+{i_}{r}-{m_}{r}"
        ws[f"{p_}{r}"] = (f"={m_}{r}+MAX(0,MIN({n_}{r},'Debt Payoff'!$C$12-SUM($W{r}:$AF{r})"
                          f"-SUMIFS($AG{r}:$AP{r},$AG$4:$AP$4,\"<\"&{n_}$4)))")
        ws[f"{b}{r}"] = f"=MAX(0,ROUND({b}{r-1}+{i_}{r}-{p_}{r},2))"
    ws.cell(r, 53).value = f"=SUM(C{r}:L{r})"
    ws.cell(r, 54).value = f"=SUM(AQ{r}:AZ{r})"
    ws.cell(r, 55).value = f"=SUM(M{r}:V{r})"
    for c in range(3, 56):
        ws.cell(r, c).number_format = MONEY0
        ws.cell(r, c).font = font(TEXT, 8)
ws.freeze_panes = "C6"

# ======================================================================== Instructions
ws = sheets["Instructions"]
setup_sheet(ws, {"A": 3, "B": 4, "C": 110}, tab=LAVENDER, rows=45)
merge_put(ws, 2, 2, 2, 3, "ULTIMATE BUDGET PLANNER", font(TEXT, 24, True), BG, align=LEFT)
put(ws, 3, 3, "Annual · Monthly · Paycheck / any period  —  dark edition", font(TEAL, 11, True))
steps = [
    ("SETUP", "Enter your budget start date, opening balance, accounts, variable-expense budgets and every "
              "recurring income, bill, debt payment, subscription and savings transfer ONCE (sample data included — "
              "overwrite or delete it). Yellow cells are inputs everywhere."),
    ("PAYMENTS", "Every recurring item auto-repeats here for 12 months, sorted by date and coloured by category. "
                 "Mark ✓ paid / ✗ skipped (or leave 'Auto-mark paid' = Yes on Setup). Enter one-off changes in the "
                 "VARIATION columns (new date, amount or account). Use the filter arrows to sort / filter."),
    ("MANUAL LOG", "Log variable spending, irregular income, extra bills and transfers between accounts. "
                   "Tip: sub-categories must match the names in Setup to land in the right table."),
    ("ACCOUNTS", "Live balance of each account and your net worth."),
    ("DASHBOARD", "Switch between Weekly / Bi-weekly / 4-weekly (paycheck) / Monthly / Quarterly / 6-month / Annual / "
                  "Custom and pick a start date — the whole dashboard recalculates for that period."),
    ("JAN – DEC", "12 monthly dashboards: budget vs actual progress, summary, charts, top expenses and daily "
                  "spending. Each month's start balance rolls over from the previous month (or type an override "
                  "for a zero-based budget). Tabs follow the start date on Setup — keep it on 1 January for the "
                  "tab names to match."),
    ("ANNUAL", "Year view of every category and sub-category by month with totals, averages and % of budget."),
    ("CALENDAR", "Pick any month: scheduled items per day, projected daily balance and weekly balance summary. "
                 "Paid items turn grey."),
    ("SAVINGS", "Up to 20 sinking funds — progress, amount left and how much to save each month."),
    ("DEBT PAYOFF", "Up to 10 debts with Snowball, Avalanche, Custom or Credit-Score-Focus ordering, "
                    "payoff dates, total interest and a month-by-month schedule (Debt Schedule tab)."),
    ("NOTES", "Built for Microsoft Excel 365 / 2021+. Don't type over white/black formula cells. "
              "If you edit recurring items after marking payments, re-check the ✓ marks (rows re-sort by date)."),
]
r = 5
for t, txt in steps:
    put(ws, r, 3, t, font(AMBER, 11, True))
    c = put(ws, r + 1, 3, txt, font(TEXT, 10), align=WRAP)
    ws.row_dimensions[r + 1].height = 30 if len(txt) < 180 else 44
    r += 3

# tab order already matches; open on Instructions
wb.active = 0
wb.calculation.fullCalcOnLoad = True
out = "Ultimate_Budget_Planner.xlsx"
wb.save(out)
print("saved", out)
