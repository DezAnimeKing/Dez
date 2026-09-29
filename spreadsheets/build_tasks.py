"""Builds Task_Habit_Tracker.xlsx — tasks, planners, decision matrix, habits, mood, gratitude, goals."""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.workbook.defined_name import DefinedName

from common import *

D = dt.date
T = dt.time
START_DATE = D(2026, 9, 1)

RS0, RSN = 6, 40            # Recurring Schedule rows 6..45
RS1 = RS0 + RSN - 1
RL0, RLN = 6, 1500          # Recurring Log rows 6..1505
RL1 = RL0 + RLN - 1
VT0, VTN = 6, 500           # Variable Tasks rows 6..505
VT1 = VT0 + VTN - 1
CN = RLN + VTN              # combined rows in Calc: 2..2001
C0, C1 = 2, CN + 1
DAYS = 366
HAB0, HABN = 9, 15          # habit rows 9..23
HCOL0 = 5                   # first day column in Habit Tracker (E)
NDAYS = 364

TODAY = "Setup!$D$7"
INTERVAL = "Setup!$D$8"
DAYSTART = "Setup!$D$9"
WEEKSTART = "Setup!$D$10"
CATS = "Setup!$F$6:$F$15"
PEOPLE = "Setup!$H$6:$H$15"
STATUSES = "Setup!$J$6:$J$11"
PRIOS = "Setup!$L$6:$L$9"
CATCOLS = [BLUE, TEAL, PINK, AMBER, LAVENDER, CORAL, GREEN, ROSE, "FF9F5A", "C084FC"]


def C(c):
    """Combined-task column range on Calc."""
    return f"Calc!${c}${C0}:${c}${C1}"


wb = Workbook()
wb.remove(wb.active)
names = ["Instructions", "Setup", "Dashboard", "Recurring Schedule", "Variable Tasks", "Recurring Log",
         "Tasks Filter", "Daily", "Weekly", "Monthly", "Decision Matrix", "Habit Tracker", "Gratitude Log",
         "Goals", "Lists", "Calc"]
sh = {n: wb.create_sheet(n) for n in names}

# ============================================================================ Lists
ws = sh["Lists"]
for c, vals in {"A": ["Freq", "Daily", "Weekdays", "Weekly", "Bi-weekly", "Monthly", "Quarterly", "Annual", "Once"],
                "B": ["Days", 1, 0, 7, 14, 0, 0, 0, 0],
                "C": ["Months", 0, 0, 0, 0, 1, 3, 12, 0],
                "D": ["Tick", "✓"],
                "E": ["Interval", 15, 30, 60],
                "F": ["Week start", "Monday", "Sunday"],
                "G": ["Sort by", "Due date", "Priority", "Category", "Person", "Status"],
                "H": ["Order", "Ascending", "Descending"],
                "I": ["Important", "Both", "Yes", "No"],
                "J": ["Decision", "DO FIRST", "SCHEDULE", "DELEGATE", "DELETE"],
                "K": ["YesNo", "Yes", "No"]}.items():
    for i, v in enumerate(vals):
        ws[f"{c}{i+1}"] = v
ws.sheet_state = "hidden"

# ============================================================================ Setup
ws = sh["Setup"]
setup_sheet(ws, {"A": 2, "B": 2, "C": 20, "D": 14, "E": 3, "F": 16, "G": 3, "H": 16, "I": 3, "J": 16, "K": 3,
                 "L": 16}, tab=AMBER, rows=30)
merge_put(ws, 2, 3, 2, 12, "SETUP", font(TEXT, 20, True), BG, align=LEFT)
put(ws, 3, 3, "Yellow cells are inputs. Categories, people, statuses & priorities feed every dropdown.",
    font(MUTED, 9, italic=True))
band(ws, 4, 3, 4, "GENERAL", LAVENDER)
gen = [("Planner start date", START_DATE, DATE), ("Today (override)", None, DATE), ("Today used", "=IF(D6=\"\",TODAY(),D6)", DATE),
       ("Schedule interval (min)", 30, "0"), ("Day starts at", T(7, 0), "h:mm AM/PM"), ("Week starts on", "Monday", None)]
for i, (t, v, f) in enumerate(gen):
    r = 5 + i
    put(ws, r, 3, t, font(MUTED, 9), PANEL, align=LEFT)
    if i == 2:
        put(ws, r, 4, v, font(TEXT, 10, True), PANEL, f, CENTER)
    else:
        input_cell(ws, r, 4, v, f)
dv_list(ws, "=Lists!$E$2:$E$4", "D8")
dv_list(ws, "=Lists!$F$2:$F$3", "D10")
put(ws, 12, 3, "Planner start = first day of the habit tracker & recurring-task horizon (1 year).",
    font(MUTED, 8, italic=True))
setup_lists = [(6, "CATEGORIES", ["Business", "Personal", "Family", "Health", "Social", "Birthdays"], BLUE),
               (8, "PEOPLE", ["Me", "Sam", "Lucy", "Tony", "Jamie", "Rebecca"], TEAL),
               (10, "STATUS", ["Not Started", "In Progress", "Waiting", "On Hold", "Done", "Cancelled"], PINK),
               (12, "PRIORITY", ["1-Urgent", "2-High", "3-Medium", "4-Low"], AMBER)]
for c, t, vals, colr in setup_lists:
    band(ws, 4, c, c, t, colr)
    n = 4 if t == "PRIORITY" else (6 if t == "STATUS" else 10)
    for i in range(n):
        v = vals[i] if i < len(vals) else None
        if t in ("STATUS", "PRIORITY"):
            put(ws, 6 + i, c, v, font(TEXT, 10, True), PANEL, align=LEFT)
        else:
            input_cell(ws, 6 + i, c, v, align=LEFT)
put(ws, 5, 6, "colour follows row order", font(MUTED, 7, italic=True))
for i in range(10):
    cf_formula(ws, f"F{6+i}", f'F{6+i}<>""', DARK, CATCOLS[i])
put(ws, 17, 10, "Status & priority lists are fixed (formulas rely on them).", font(MUTED, 7, italic=True))

# ============================================================================ Recurring Schedule
ws = sh["Recurring Schedule"]
setup_sheet(ws, {"A": 2, "B": 4, "C": 28, "D": 13, "E": 12, "F": 11, "G": 11, "H": 10, "I": 10, "J": 11, "K": 10,
                 "L": 12, "M": 28}, tab=PINK, rows=50)
merge_put(ws, 2, 2, 2, 10, "RECURRING TASK SCHEDULE", font(TEXT, 18, True), BG, align=LEFT)
put(ws, 3, 2, "Set up repeating tasks ONCE — they auto-fill the Recurring Log, Tasks Filter, Dashboard, Daily, "
              "Weekly, Monthly & Decision Matrix for a full year.", font(MUTED, 9, italic=True))
band(ws, 4, 2, 13, "RECURRING TASKS", PINK)
for i, h in enumerate(["#", "TASK", "CATEGORY", "FREQUENCY", "1ST DATE", "END DATE", "START TIME", "END TIME",
                       "PRIORITY", "IMPORTANT", "PERSON", "NOTES"]):
    put(ws, 5, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
rec = [("Yoga", "Personal", "Daily", D(2026, 9, 28), None, T(7, 30), T(8, 0), "3-Medium", "✓", "Me"),
       ("Meditate", "Health", "Weekdays", D(2026, 9, 28), None, T(8, 30), T(8, 45), "4-Low", "", "Me"),
       ("Monday Meeting", "Business", "Weekly", D(2026, 9, 28), None, T(9, 0), T(10, 0), "2-High", "✓", "Me"),
       ("Write project proposal", "Business", "Weekly", D(2026, 9, 30), D(2026, 12, 31), T(9, 15), T(11, 0),
        "1-Urgent", "✓", "Me"),
       ("Pilates", "Health", "Weekly", D(2026, 10, 1), None, T(18, 0), T(19, 0), "3-Medium", "✓", "Me"),
       ("Amy's Soccer Game", "Family", "Weekly", D(2026, 10, 3), D(2026, 12, 19), T(10, 0), T(11, 30), "2-High",
        "✓", "Sam"),
       ("Water Plants", "Personal", "Weekly", D(2026, 10, 4), None, T(17, 0), T(17, 15), "4-Low", "", "Lucy"),
       ("Friday task review", "Business", "Weekly", D(2026, 10, 2), None, T(16, 0), T(16, 30), "3-Medium", "",
        "Me"),
       ("Date Night", "Family", "Bi-weekly", D(2026, 10, 3), None, T(19, 0), T(21, 0), "2-High", "✓", "Me"),
       ("Monthly goal planning", "Personal", "Monthly", D(2026, 10, 1), None, T(20, 0), T(21, 0), "2-High", "✓",
        "Me"),
       ("Pay credit card", "Personal", "Monthly", D(2026, 10, 10), None, None, None, "1-Urgent", "✓", "Me"),
       ("Dentist check-up", "Health", "Quarterly", D(2026, 11, 12), None, T(14, 0), T(15, 0), "3-Medium", "✓",
        "Me")]
for i in range(RSN):
    r = RS0 + i
    v = rec[i] if i < len(rec) else (None,) * 10
    put(ws, r, 2, i + 1, font(MUTED, 8), PANEL, align=CENTER)
    fm = [None, None, None, DATE, DATE, "h:mm AM/PM", "h:mm AM/PM", None, None, None]
    for j in range(10):
        input_cell(ws, r, 3 + j, v[j], fm[j], LEFT if j == 0 else CENTER)
    input_cell(ws, r, 13, None, align=LEFT)
    ws[f"N{r}"] = f'=IFERROR(VLOOKUP(E{r},Lists!$A$2:$C$9,2,0),0)'
    ws[f"O{r}"] = f'=IFERROR(VLOOKUP(E{r},Lists!$A$2:$C$9,3,0),0)'
    ws[f"P{r}"] = f'=IF(G{r}="",Setup!$D$5+{DAYS-1},MIN(G{r},Setup!$D$5+{DAYS-1}))'
for c in "NOP":
    ws.column_dimensions[c].hidden = True
dv_list(ws, f"={CATS}", f"D{RS0}:D{RS1}")
dv_list(ws, "=Lists!$A$2:$A$9", f"E{RS0}:E{RS1}")
dv_list(ws, f"={PRIOS}", f"J{RS0}:J{RS1}")
dv_list(ws, "=Lists!$D$2", f"K{RS0}:K{RS1}")
dv_list(ws, f"={PEOPLE}", f"L{RS0}:L{RS1}")
ws.freeze_panes = "C6"

# ============================================================================ Calc
ws = sh["Calc"]
ws["A1"] = "date"
for j in range(RSN):
    ws.cell(1, 2 + j, j + 1)
ws["AP1"] = "count"; ws["AQ1"] = "before"
for d in range(DAYS):
    r = 2 + d
    ws[f"A{r}"] = f"=Setup!$D$5+{d}"
    for j in range(RSN):
        sr = RS0 + j
        rs = "'Recurring Schedule'!"
        F, E_, N_, O_, P_ = (f"{rs}${c}${sr}" for c in "FENOP")
        test = (f"IF({N_}>0,MOD($A{r}-{F},{N_})=0,IF({O_}>0,AND(DAY($A{r})=MIN(DAY({F}),DAY(EOMONTH($A{r},0))),"
                f"MOD((YEAR($A{r})-YEAR({F}))*12+MONTH($A{r})-MONTH({F}),{O_})=0),"
                f"IF({E_}=\"Weekdays\",WEEKDAY($A{r},2)<=5,$A{r}={F})))")
        ws.cell(r, 2 + j).value = (f'=IF({rs}$C${sr}="","",IF({F}="","",IF(OR($A{r}<{F},$A{r}>{P_}),"",'
                                   f'IF({test},{j+1},""))))')
    ws[f"AP{r}"] = f"=COUNT(B{r}:AO{r})"
    ws[f"AQ{r}"] = "=0" if d == 0 else f"=AQ{r-1}+AP{r-1}"
ws[f"AQ{DAYS+2}"] = f"=AQ{DAYS+1}+AP{DAYS+1}"

# combined task table AS..BN
heads = {"AS": "src", "AT": "row", "AU": "task", "AV": "cat", "AW": "due", "AX": "start", "AY": "end",
         "AZ": "status", "BA": "prio", "BB": "imp", "BC": "person", "BD": "decision", "BE": "progress",
         "BF": "prionum", "BG": "done", "BH": "valid", "BI": "include", "BJ": "sortkey", "BK": "fkey",
         "BL": "overdue", "BM": "week", "BN": "pkey"}
for c, h in heads.items():
    ws[f"{c}1"] = h
TF = "'Tasks Filter'!"
for i in range(1, CN + 1):
    r = i + 1
    if i <= RLN:
        s, sr = "'Recurring Log'!", RL0 + i - 1
        ws[f"AS{r}"] = "R"
    else:
        s, sr = "'Variable Tasks'!", VT0 + i - RLN - 1
        ws[f"AS{r}"] = "V"
    ws[f"AT{r}"] = sr
    ws[f"AU{r}"] = f'={s}C{sr}&""'
    ws[f"AV{r}"] = f'={s}D{sr}&""'
    ws[f"AW{r}"] = f'=IF({s}E{sr}="","",{s}E{sr})'
    ws[f"AX{r}"] = f'=IF({s}G{sr}="","",{s}G{sr})'
    ws[f"AY{r}"] = f'=IF({s}H{sr}="","",{s}H{sr})'
    ws[f"AZ{r}"] = f'=IF(AU{r}="","",IF({s}J{sr}="","Not Started",{s}J{sr}))'
    ws[f"BA{r}"] = f'={s}L{sr}&""'
    ws[f"BB{r}"] = f'={s}M{sr}&""'
    ws[f"BC{r}"] = f'={s}N{sr}&""'
    ws[f"BD{r}"] = f'={s}O{sr}&""'
    ws[f"BE{r}"] = f'=IF({s}K{sr}="","",{s}K{sr})'
    ws[f"BF{r}"] = f"=IFERROR(MATCH(BA{r},{PRIOS},0),5)"
    ws[f"BG{r}"] = f'=IF(OR(AZ{r}="Done",AZ{r}="Cancelled"),1,0)'
    ws[f"BH{r}"] = f'=IF(AND(AU{r}<>"",ISNUMBER(AW{r})),1,0)'
    ws[f"BI{r}"] = (f'=IF(BH{r}=0,0,IF(AND(AW{r}>={TF}$F$3,AW{r}<={TF}$F$4,'
                    f'OR({TF}$I$3="",ISNUMBER(SEARCH({TF}$I$3,AU{r}))),'
                    f'OR({TF}$I$4="",BF{r}=IFERROR(MATCH({TF}$I$4,{PRIOS},0),0)),'
                    f'OR({TF}$I$5="Both",AND({TF}$I$5="Yes",BB{r}="✓"),AND({TF}$I$5="No",BB{r}<>"✓")),'
                    f'OR({TF}$L$3="",AV{r}={TF}$L$3),OR({TF}$L$4="",AZ{r}={TF}$L$4),'
                    f'OR({TF}$L$5="",BC{r}={TF}$L$5),OR({TF}$O$3="",BD{r}={TF}$O$3),'
                    f'OR({TF}$O$4="Yes",BG{r}=0)),1,0))')
    ws[f"BJ{r}"] = (f'=IF(BH{r}=0,"",IF({TF}$F$6="Descending",-1,1)*CHOOSE(MATCH({TF}$F$5,Lists!$G$2:$G$6,0),'
                    f'AW{r},BF{r},IFERROR(MATCH(AV{r},{CATS},0),99),IFERROR(MATCH(BC{r},{PEOPLE},0),99),'
                    f'IFERROR(MATCH(AZ{r},{STATUSES},0),99))*1E+9+(AW{r}-40000)*10000+{i})')
    ws[f"BK{r}"] = f'=IF(BI{r}=1,BJ{r},"")'
    ws[f"BL{r}"] = f'=IF(AND(BH{r}=1,BG{r}=0,N(AW{r})<{TODAY}),(AW{r}-40000)*10000+{i},"")'
    ws[f"BM{r}"] = (f'=IF(AND(BH{r}=1,AZ{r}<>"Cancelled",N(AW{r})>=Dashboard!$AB$3,N(AW{r})<=Dashboard!$AB$3+6),'
                    f'(AW{r}-40000)*100000+BF{r}*10000+{i},"")')
    ws[f"BN{r}"] = f'=IF(AND(BH{r}=1,AZ{r}<>"Cancelled"),BF{r}*10000+{i},"")'
ws.sheet_state = "hidden"
for nm, c in [("TaskCat", "AV"), ("TaskDone", "BG"), ("TaskDue", "AW"), ("TaskStat", "AZ")]:
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"Calc!${c}${C0}:${c}${C1}")


# ============================================================================ task list sheets
LIST_HEADS = ["#", "TASK", "CATEGORY", "DUE DATE", "DAYS LEFT", "START", "END", "DURATION", "STATUS",
              "PROGRESS", "PRIORITY", "IMPORTANT", "DELEGATE TO", "DECISION", "NOTES"]
LIST_W = {"A": 2, "B": 5, "C": 28, "D": 12, "E": 11, "F": 9, "G": 10, "H": 10, "I": 12, "J": 13, "K": 10,
          "L": 11, "M": 10, "N": 13, "O": 12, "P": 28}


def derived(ws, r):
    ws[f"F{r}"] = f'=IF(OR(C{r}="",E{r}=""),"",E{r}-{TODAY})'
    ws[f"I{r}"] = (f'=IF(OR(G{r}="",H{r}=""),"",TRIM(IF(HOUR(H{r}-G{r})>0,HOUR(H{r}-G{r})&" hr ","")&'
                   f'IF(MINUTE(H{r}-G{r})>0,MINUTE(H{r}-G{r})&" min","")))')
    ws[f"O{r}"] = (f'=IF(C{r}="","",IF(AND(OR(L{r}="1-Urgent",L{r}="2-High"),M{r}="✓"),"DO FIRST",'
                   f'IF(M{r}="✓","SCHEDULE",IF(OR(L{r}="1-Urgent",L{r}="2-High"),"DELEGATE","DELETE"))))')


def list_cf(ws, r0, r1):
    rng = f"B{r0}:P{r1}"
    cf_formula(ws, rng, f'OR($J{r0}="Done",$J{r0}="Cancelled")', "6E6E86", stop=True)
    cf_formula(ws, rng, f'AND($C{r0}<>"",ISNUMBER($E{r0}),$E{r0}<{TODAY})', DARK, "FF5C9E")
    cf_formula(ws, rng, f'AND($C{r0}<>"",$E{r0}={TODAY})', TEXT, "3B2F63")
    for i in range(6):
        cf_formula(ws, f"D{r0}:D{r1}", f'AND($D{r0}<>"",$D{r0}=INDEX({CATS},{i+1}))', CATCOLS[i])
    for t, c in (("DO FIRST", CORAL), ("SCHEDULE", TEAL), ("DELEGATE", AMBER), ("DELETE", LAVENDER)):
        cf_formula(ws, f"O{r0}:O{r1}", f'$O{r0}="{t}"', c, bold=True)
    databar(ws, f"K{r0}:K{r1}", TEAL)


def list_header(ws, title, sub, colr):
    setup_sheet(ws, LIST_W, tab=colr, zoom=85)
    merge_put(ws, 2, 2, 2, 10, title, font(TEXT, 18, True), BG, align=LEFT)
    put(ws, 3, 2, sub, font(MUTED, 9, italic=True))
    band(ws, 4, 2, 4, "TASK DETAILS", colr)
    band(ws, 4, 5, 9, "SCHEDULE", LAVENDER)
    band(ws, 4, 10, 11, "STATUS & PROGRESS", TEAL)
    band(ws, 4, 12, 15, "EISENHOWER MATRIX", AMBER)
    band(ws, 4, 16, 16, "NOTES", "5C5C78", text_color=TEXT)
    for i, h in enumerate(LIST_HEADS):
        put(ws, 5, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)


def style_row(ws, r, alt, inputs):
    for c in range(2, 17):
        cell = ws.cell(r, c)
        cell.font = font(TEXT, 9)
        cell.fill = fill("2E2A1E" if c in inputs else (ROW if alt else ROW2))
        cell.alignment = LEFT if c in (3, 16) else CENTER
        cell.border = BOTTOM
    ws.cell(r, 5).number_format = DATE
    ws.cell(r, 6).number_format = '0;-0;0'
    for c in (7, 8):
        ws.cell(r, c).number_format = "h:mm AM/PM"
    ws.cell(r, 11).number_format = "0%"


def list_dv(ws, r0, r1, full):
    dv_list(ws, f"={STATUSES}", f"J{r0}:J{r1}")
    if full:
        dv_list(ws, f"={CATS}", f"D{r0}:D{r1}")
        dv_list(ws, f"={PRIOS}", f"L{r0}:L{r1}")
        dv_list(ws, "=Lists!$D$2", f"M{r0}:M{r1}")
        dv_list(ws, f"={PEOPLE}", f"N{r0}:N{r1}")


# ---- Variable Tasks
ws = sh["Variable Tasks"]
list_header(ws, "VARIABLE TASKS", "One-off tasks — type them here. Pink = overdue · purple = due today · grey = done.",
            BLUE)
random.seed(3)
vt_names = [("Respond to emails", "Business"), ("Conduct market research", "Business"),
            ("Organize project files", "Business"), ("Brief Amanda", "Business"), ("Track weekly goals", "Personal"),
            ("Coffee with Alex", "Social"), ("Declutter desk", "Personal"), ("Review client feedback", "Business"),
            ("Analyze sales data", "Business"), ("Take Sam to Dentist", "Family"), ("Thyroid Check", "Health"),
            ("Design product logo", "Business"), ("Draft blog post", "Business"), ("Dinner with Amy", "Social"),
            ("Plan weekend away", "Family"), ("Clean kitchen", "Family"), ("Family Lunch", "Family"),
            ("Madison's Bday", "Birthdays"), ("Client follow-up calls", "Business"),
            ("Schedule team meeting", "Business"), ("Plan marketing campaign", "Business"),
            ("Review ad performance", "Business"), ("Presentation Project X", "Business"), ("Mole Check", "Health"),
            ("Running", "Personal"), ("Doctor Check Up", "Health"), ("Buy groceries", "Family"),
            ("John's Birthday", "Birthdays"), ("Test product features", "Business"), ("Girls lunch", "Social"),
            ("Facial Treatment", "Health"), ("Book massage", "Personal"), ("Hair dresser", "Personal"),
            ("Finalize budget allocation", "Business"), ("Call client", "Business"), ("Prepare launch checklist", "Business"),
            ("Email client re: project X", "Business"), ("Organize photos", "Personal"), ("Car service", "Family"),
            ("Renew passport", "Personal")]
people = ["Me", "Me", "Sam", "Lucy", "Tony", "Jamie", "Rebecca", "Me"]
prios = ["1-Urgent", "2-High", "3-Medium", "4-Low"]
vts = []
for k, (nm, cat) in enumerate(vt_names):
    due = D(2026, 9, 20) + dt.timedelta(days=int(k * 1.3))
    hr = random.choice([None, 9, 10, 11, 13, 15, 16])
    st = T(hr, random.choice([0, 30])) if hr else None
    en = T(hr + random.choice([0, 1, 1, 2]), 45 if st and st.minute == 0 else 30) if hr else None
    if due < D(2026, 9, 29):
        status = random.choice(["Done", "Done", "Done", "In Progress", "Not Started"])
    else:
        status = random.choice(["Not Started", "Not Started", "In Progress", "Waiting", None])
    prog = {"Done": 1, "In Progress": random.choice([0.25, 0.4, 0.6, 0.8]), "Waiting": 0.5}.get(status)
    vts.append((nm, cat, due, st, en, status, prog, random.choice(prios), random.choice(["✓", "✓", ""]),
                random.choice(people), ""))
for i in range(VTN):
    r = VT0 + i
    ws[f"B{r}"] = i + 1
    v = vts[i] if i < len(vts) else (None,) * 11
    for c, val in zip("CDEGHJKLMNP", v):
        ws[f"{c}{r}"] = val
    derived(ws, r)
    style_row(ws, r, i % 2, {3, 4, 5, 7, 8, 10, 11, 12, 13, 14, 16})
    ws[f"B{r}"].font = font(MUTED, 8)
list_cf(ws, VT0, VT1)
list_dv(ws, VT0, VT1, True)
ws.auto_filter.ref = f"B5:P{VT1}"
ws.freeze_panes = "D6"

# ---- Recurring Log
ws = sh["Recurring Log"]
list_header(ws, "RECURRING LOG", "Auto-generated from the Recurring Schedule — only update STATUS, PROGRESS & NOTES here.",
            PINK)
RSs = "'Recurring Schedule'!"
for i in range(RLN):
    r = RL0 + i
    k = i + 1
    ws[f"B{r}"] = f'=IF(R{r}="","",{k})'
    ws[f"Q{r}"] = f'=IF({k}>Calc!$AQ${DAYS+2},"",MATCH({k-1},Calc!$AQ$2:$AQ${DAYS+1},1))'
    ws[f"R{r}"] = (f'=IF(Q{r}="","",SMALL(INDEX(Calc!$B$2:$AO${DAYS+1},Q{r},0),'
                   f'{k}-INDEX(Calc!$AQ$2:$AQ${DAYS+1},Q{r})))')
    for c, sc in zip("CDGHLMN", "CDHIJKL"):
        ws[f"{c}{r}"] = f'=IF(R{r}="","",IF(INDEX({RSs}${sc}${RS0}:${sc}${RS1},R{r})="","",INDEX({RSs}${sc}${RS0}:${sc}${RS1},R{r})))'
    ws[f"E{r}"] = f'=IF(Q{r}="","",INDEX(Calc!$A$2:$A${DAYS+1},Q{r}))'
    derived(ws, r)
    style_row(ws, r, i % 2, {10, 11, 16})
    ws[f"B{r}"].font = font(MUTED, 8)
for c in "QR":
    ws.column_dimensions[c].hidden = True
list_cf(ws, RL0, RL1)
list_dv(ws, RL0, RL1, False)
ws.auto_filter.ref = f"B5:P{RL1}"
ws.freeze_panes = "D6"

# ============================================================================ Tasks Filter
ws = sh["Tasks Filter"]
setup_sheet(ws, dict(LIST_W, Q=9), tab=TEAL, zoom=85)
ws.row_dimensions[2].height = 30
merge_put(ws, 2, 2, 3, 4, "Tasks Filter", font(TEXT, 22, True), PANEL, align=LEFT)
put(ws, 4, 2, None, fill_color=TEAL)
ws.merge_cells("B4:D4")
put(ws, 4, 2, f'=TEXT({TODAY},"dddd, mmmm d, yyyy")', font(DARK, 9, True), TEAL, align=CENTER)
labels = [(3, 5, "From"), (4, 5, "To"), (5, 5, "Sort by"), (6, 5, "Order"),
          (3, 8, "Task name"), (4, 8, "Priority"), (5, 8, "Important"),
          (3, 11, "Category"), (4, 11, "Status"), (5, 11, "Person"),
          (3, 14, "Decision"), (4, 14, "Show done")]
vals = {(3, 6): (f"={TODAY}-30", DATE), (4, 6): (f"={TODAY}+90", DATE), (5, 6): ("Due date", None),
        (6, 6): ("Ascending", None), (3, 9): (None, None), (4, 9): (None, None), (5, 9): ("Both", None),
        (3, 12): (None, None), (4, 12): (None, None), (5, 12): (None, None), (3, 15): (None, None),
        (4, 15): ("Yes", None)}
band(ws, 2, 5, 15, "DATE RANGE · SORT · FILTER  (blank = all)", LAVENDER, 9)
for r, c, t in labels:
    put(ws, r, c, t, font(MUTED, 8, True), PANEL, align=RIGHT)
for (r, c), (v, f) in vals.items():
    input_cell(ws, r, c, v, f)
dv_list(ws, "=Lists!$G$2:$G$6", "F5"); dv_list(ws, "=Lists!$H$2:$H$3", "F6")
dv_list(ws, f"={PRIOS}", "I4"); dv_list(ws, "=Lists!$I$2:$I$4", "I5")
dv_list(ws, f"={CATS}", "L3"); dv_list(ws, f"={STATUSES}", "L4"); dv_list(ws, f"={PEOPLE}", "L5")
dv_list(ws, "=Lists!$J$2:$J$5", "O3"); dv_list(ws, "=Lists!$K$2:$K$3", "O4")
for c, (t, f, colr) in zip((16, 17), [("SHOWN", "=COUNT(Calc!$BK$2:$BK$2001)", TEAL),
                                      ("OVERDUE", "=COUNT(Calc!$BL$2:$BL$2001)", CORAL)]):
    put(ws, 3, c, t, font(MUTED, 8, True), PANEL, align=CENTER)
    merge_put(ws, 4, c, 5, c, f, font(colr, 18, True), PANEL, "0")
put(ws, 6, 16, "DUE TODAY", font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 7, 16, f"=COUNTIFS(Calc!$AW$2:$AW$2001,{TODAY},Calc!$BH$2:$BH$2001,1)", font(AMBER, 14, True), PANEL,
    "0", CENTER)
put(ws, 6, 17, "DONE", font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 7, 17, "=SUMPRODUCT((Calc!$BI$2:$BI$2001=1)*(Calc!$BG$2:$BG$2001=1))", font(LAVENDER, 14, True),
    PANEL, "0", CENTER)
band(ws, 8, 2, 4, "TASK DETAILS", TEAL)
band(ws, 8, 5, 9, "SCHEDULE", LAVENDER)
band(ws, 8, 10, 11, "STATUS & PROGRESS", TEAL)
band(ws, 8, 12, 15, "EISENHOWER MATRIX", AMBER)
band(ws, 8, 16, 16, "EDIT", "5C5C78", text_color=TEXT)
heads = LIST_HEADS[:-1] + ["EDIT TASK"]
for i, h in enumerate(heads):
    put(ws, 9, 2 + i, h, font(MUTED, 8, True), PANEL, align=CENTER, border=BOTTOM)
CM = {"C": "AU", "D": "AV", "E": "AW", "G": "AX", "H": "AY", "J": "AZ", "K": "BE", "L": "BA", "M": "BB",
      "N": "BC", "O": "BD"}
FV0, FVN = 10, 300
for k in range(1, FVN + 1):
    r = FV0 + k - 1
    ws[f"R{r}"] = f'=IFERROR(MATCH(SMALL(Calc!$BK$2:$BK$2001,{k}),Calc!$BK$2:$BK$2001,0),"")'
    ws[f"B{r}"] = f'=IF(R{r}="","",{k})'
    for c, cc in CM.items():
        ws[f"{c}{r}"] = f'=IF($R{r}="","",IF(INDEX(Calc!${cc}$2:${cc}$2001,$R{r})="","",INDEX(Calc!${cc}$2:${cc}$2001,$R{r})))'
    ws[f"F{r}"] = f'=IF(OR(C{r}="",E{r}=""),"",E{r}-{TODAY})'
    ws[f"I{r}"] = (f'=IF(OR(G{r}="",H{r}=""),"",TRIM(IF(HOUR(H{r}-G{r})>0,HOUR(H{r}-G{r})&" hr ","")&'
                   f'IF(MINUTE(H{r}-G{r})>0,MINUTE(H{r}-G{r})&" min","")))')
    ws[f"P{r}"] = (f'=IF(R{r}="","",HYPERLINK("#\'"&IF(INDEX(Calc!$AS$2:$AS$2001,R{r})="R","Recurring Log",'
                   f'"Variable Tasks")&"\'!J"&INDEX(Calc!$AT$2:$AT$2001,R{r}),"✎ edit"))')
    style_row(ws, r, k % 2, set())
    ws[f"P{r}"].font = font(TEAL, 9, True)
ws.column_dimensions["R"].hidden = True
list_cf(ws, FV0, FV0 + FVN - 1)
ws.freeze_panes = "D10"

# ============================================================================ Dashboard
ws = sh["Dashboard"]
setup_sheet(ws, {"A": 2, "B": 24, "C": 2, "D": 4, "E": 26, "F": 12, "G": 11, "H": 3, "I": 2, "J": 4, "K": 26,
                 "L": 12, "M": 11, "N": 3, "O": 2, "P": 14, "Q": 10, "R": 10, "S": 10, "T": 10, "U": 10},
            tab=LAVENDER, zoom=85, rows=62)
paint(ws, 2, 2, 60, 2, PANEL)
put(ws, 2, 2, "TODAY IS", font(MUTED, 9, True), PANEL, align=CENTER)
put(ws, 3, 2, f"={TODAY}", font(TEXT, 13, True), PANEL, 'mmmm d, yyyy', CENTER)
nav = [("SETUP", "Setup"), None, ("TASK LOG + FILTER", None), ("Recurring Schedule", "Recurring Schedule"),
       ("Variable Tasks", "Variable Tasks"), ("Recurring Log", "Recurring Log"), ("Tasks Filter", "Tasks Filter"),
       ("Gratitude Log", "Gratitude Log"), ("Habit Tracker", "Habit Tracker"), None, ("DASHBOARDS", None),
       ("Daily", "Daily"), ("Weekly", "Weekly"), ("Monthly", "Monthly"), ("Decision Matrix", "Decision Matrix"),
       None, ("BONUS", None), ("Goals", "Goals"), ("Instructions", "Instructions")]
r = 5
for item in nav:
    if item:
        t, target = item
        if target:
            put(ws, r, 2, f'=HYPERLINK("#\'{target}\'!A1","›  {t}")', font(TEAL, 10, True), PANEL, align=LEFT)
        else:
            put(ws, r, 2, t, font(DARK, 9, True), LAVENDER, align=LEFT)
    r += 1
band(ws, 26, 2, 2, "MY TOP HABITS", AMBER, 9)
for i in range(8):
    put(ws, 27 + i, 2, f"=IF('Habit Tracker'!B{HAB0+i}=\"\",\"\",\"{i+1}.  \"&'Habit Tracker'!B{HAB0+i})",
        font(TEXT, 9), PANEL, align=LEFT)

# today's tasks
band(ws, 2, 4, 6, "TODAY'S TASKS", TEAL)
put(ws, 2, 7, f'=IFERROR(COUNTIFS(TaskDue,{TODAY},TaskDone,1)/COUNTIFS(TaskDue,{TODAY},TaskStat,"<>Cancelled"),0)',
    font(DARK, 12, True), TEAL, PCT, CENTER)
for i, h in enumerate(["", "TASK", "CATEGORY", "PRIORITY"]):
    put(ws, 3, 4 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i != 1 else LEFT)


def task_rows(ws, r0, n, key_expr, cols, idxcol, extra=None):
    """cols: list of (col letter, calc column or special)."""
    for k in range(1, n + 1):
        r = r0 + k - 1
        ws[f"{idxcol}{r}"] = key_expr(k)
        for c, src in cols:
            if src == "tick":
                v = f'=IF({idxcol}{r}="","",IF(INDEX(Calc!$BG$2:$BG$2001,{idxcol}{r})=1,"✓","○"))'
            elif src == "late":
                v = f'=IF({idxcol}{r}="","",{TODAY}-INDEX(Calc!$AW$2:$AW$2001,{idxcol}{r}))'
            else:
                v = f'=IF({idxcol}{r}="","",INDEX(Calc!${src}$2:${src}$2001,{idxcol}{r})&"")'
                if src == "AW":
                    v = f'=IF({idxcol}{r}="","",INDEX(Calc!$AW$2:$AW$2001,{idxcol}{r}))'
            put(ws, r, ord(c) - 64 if len(c) == 1 else None, v, font(TEXT, 9), PANEL if k % 2 else ROW2,
                DATE if src == "AW" else ("0" if src == "late" else None),
                LEFT if src in ("AU",) else CENTER, BOTTOM)
        rng = f"{cols[0][0]}{r}:{cols[-1][0]}{r}"
        cf_formula(ws, rng, f'AND(${idxcol}{r}<>"",INDEX(TaskDone,${idxcol}{r})=1)', "6E6E86")
    ws.column_dimensions[idxcol].hidden = True


task_rows(ws, 4, 14, lambda k: f'=IFERROR(MOD(_xlfn.AGGREGATE(15,6,Calc!$BN$2:$BN$2001/(Calc!$AW$2:$AW$2001={TODAY}),{k}),10000),"")',
          [("D", "tick"), ("E", "AU"), ("F", "AV"), ("G", "BA")], "H")
band(ws, 20, 4, 6, "OVERDUE TASKS", CORAL)
put(ws, 20, 7, "=COUNT(Calc!$BL$2:$BL$2001)", font(DARK, 12, True), CORAL, "0", CENTER)
for i, h in enumerate(["", "TASK", "DAYS LATE", "PRIORITY"]):
    put(ws, 21, 4 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i != 1 else LEFT)
task_rows(ws, 22, 12, lambda k: f'=IFERROR(MATCH(SMALL(Calc!$BL$2:$BL$2001,{k}),Calc!$BL$2:$BL$2001,0),"")',
          [("D", "tick"), ("E", "AU"), ("F", "late"), ("G", "BA")], "I")
# this week
band(ws, 2, 10, 12, "THIS WEEK'S TASKS", AMBER)
ws["AB3"] = f'={TODAY}-MOD(WEEKDAY({TODAY},IF({WEEKSTART}="Monday",2,1))-1,7)'
ws["AB3"].number_format = DATE
put(ws, 2, 13, f'=IFERROR(COUNTIFS(TaskDue,">="&$AB$3,TaskDue,"<="&($AB$3+6),TaskDone,1)/'
               f'COUNTIFS(TaskDue,">="&$AB$3,TaskDue,"<="&($AB$3+6),TaskStat,"<>Cancelled"),0)', font(DARK, 12, True),
    AMBER, PCT, CENTER)
put(ws, 3, 11, '="Week starting "&TEXT($AB$3,"ddd d mmm")', font(MUTED, 8, True), PANEL, align=LEFT)
for i, h in enumerate(["", "", "CATEGORY", "DUE"]):
    put(ws, 3, 10 + i, h if i != 1 else None, font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 3, 11, '="TASK  ·  week of "&TEXT($AB$3,"d mmm")', font(MUTED, 8, True), PANEL, align=LEFT)
task_rows(ws, 4, 30, lambda k: f'=IFERROR(MATCH(SMALL(Calc!$BM$2:$BM$2001,{k}),Calc!$BM$2:$BM$2001,0),"")',
          [("J", "tick"), ("K", "AU"), ("L", "AV"), ("M", "AW")], "N")
for c in ("E", "K"):
    for i in range(6):
        cf_formula(ws, f"{'F' if c == 'E' else 'L'}4:{'F' if c == 'E' else 'L'}33",
                   f'AND({"F" if c == "E" else "L"}4<>"",{"F" if c == "E" else "L"}4=INDEX({CATS},{i+1}))', CATCOLS[i])
# summary with filters
band(ws, 2, 16, 21, "TASKS SUMMARY", LAVENDER)
for i, (t, v, f) in enumerate([("Start date", f"=DATE(YEAR({TODAY}),MONTH({TODAY}),1)", DATE),
                               ("End date", f"=EOMONTH({TODAY},0)", DATE), ("Category", None, None),
                               ("Person", None, None)]):
    put(ws, 3 + i, 16, t, font(MUTED, 9), PANEL, align=LEFT)
    ws.merge_cells(start_row=3 + i, start_column=17, end_row=3 + i, end_column=18)
    input_cell(ws, 3 + i, 17, v, f)
    ws.cell(3 + i, 18).fill = fill(INPUT)
dv_list(ws, f"={CATS}", "Q5"); dv_list(ws, f"={PEOPLE}", "Q6")
put(ws, 3, 19, "blank = all", font(MUTED, 8, italic=True), PANEL)
flt = (f'Calc!$AW$2:$AW$2001,">="&$Q$3,Calc!$AW$2:$AW$2001,"<="&$R$4*0+$Q$4,Calc!$BH$2:$BH$2001,1,'
       f'Calc!$AV$2:$AV$2001,IF($Q$5="","*",$Q$5),Calc!$BC$2:$BC$2001,IF($Q$6="","*",$Q$6)')
flt = flt.replace('"<="&$R$4*0+$Q$4', '"<="&$Q$4')
# helper tables (visible, compact) at P8..
band(ws, 8, 16, 18, "DECISION MATRIX", CORAL, 9)
put(ws, 9, 16, "DECISION", font(MUTED, 8, True), PANEL, align=LEFT)
put(ws, 9, 17, "TO DO", font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 9, 18, "DONE", font(MUTED, 8, True), PANEL, align=CENTER)
for i, (t, colr) in enumerate([("DO FIRST", CORAL), ("SCHEDULE", TEAL), ("DELEGATE", AMBER), ("DELETE", LAVENDER)]):
    r = 10 + i
    put(ws, r, 16, t, font(colr, 9, True), PANEL, align=LEFT)
    put(ws, r, 17, f'=COUNTIFS({flt},Calc!$BD$2:$BD$2001,"{t}",Calc!$BG$2:$BG$2001,0)', font(TEXT, 10, True),
        PANEL, "0", CENTER)
    put(ws, r, 18, f'=COUNTIFS({flt},Calc!$BD$2:$BD$2001,"{t}",Calc!$BG$2:$BG$2001,1)', font(MUTED, 9), PANEL,
        "0", CENTER)
band(ws, 8, 19, 21, "BY STATUS", PINK, 9)
for i in range(6):
    r = 9 + i
    put(ws, r, 19, f"=INDEX({STATUSES},{i+1})", font(TEXT, 9), PANEL, align=LEFT)
    ws.merge_cells(start_row=r, start_column=19, end_row=r, end_column=20)
    put(ws, r, 21, f'=COUNTIFS({flt},Calc!$AZ$2:$AZ$2001,S{r})', font(TEXT, 10, True), PANEL, "0", CENTER)
# category / person counts (helper visible under charts)
for (c0, title, lst, cc) in [(23, "CATEGORY", CATS, "AV"), (25, "PERSON", PEOPLE, "BC")]:
    ws.cell(1, c0, title)
    for i in range(10):
        ws.cell(2 + i, c0).value = f'=IF(INDEX({lst},{i+1})="","",INDEX({lst},{i+1}))'
        ws.cell(2 + i, c0 + 1).value = f'=IF({col(c0)}{2+i}="","",COUNTIFS({flt},Calc!${cc}$2:${cc}$2001,{col(c0)}{2+i}))'
for c in ("W", "X", "Y", "Z", "AA", "AB"):
    ws.column_dimensions[c].hidden = True
band(ws, 16, 16, 21, "TASKS BY CATEGORY", BLUE, 9)
dn = DoughnutChart(holeSize=50)
dn.add_data(Reference(ws, min_col=24, min_row=2, max_row=11))
dn.set_categories(Reference(ws, min_col=23, min_row=2, max_row=11))
dark_chart(dn, axes=False); color_points(dn.series[0], CATCOLS); pct_labels(dn)
dn.visible_cells_only = False
dn.width, dn.height = 12.5, 6.5
ws.add_chart(dn, "P17")
band(ws, 31, 16, 21, "TASKS BY PERSON", TEAL, 9)
dn2 = DoughnutChart(holeSize=50)
dn2.add_data(Reference(ws, min_col=26, min_row=2, max_row=11))
dn2.set_categories(Reference(ws, min_col=25, min_row=2, max_row=11))
dark_chart(dn2, axes=False); color_points(dn2.series[0], PALETTE); pct_labels(dn2)
dn2.visible_cells_only = False
dn2.width, dn2.height = 12.5, 6.5
ws.add_chart(dn2, "P32")
band(ws, 36, 4, 13, "DECISION MATRIX — TO DO vs DONE", CORAL, 9)
bc = BarChart(); bc.type = "col"; bc.grouping = "clustered"; bc.gapWidth = 60
bc.add_data(Reference(ws, min_col=17, max_col=18, min_row=9, max_row=13), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=16, min_row=10, max_row=13))
dark_chart(bc, legend_pos="t"); color_series(bc.series[0], CORAL); color_series(bc.series[1], "5C5C78")
bc.width, bc.height = 17, 6.5
ws.add_chart(bc, "D37")

# ============================================================================ Daily
ws = sh["Daily"]
setup_sheet(ws, {"A": 2, "B": 10, "C": 28, "D": 3, "E": 2, "F": 4, "G": 28, "H": 12, "I": 11, "J": 3, "K": 2,
                 "L": 34, "M": 6}, tab=TEAL, zoom=85, rows=50)
ws.row_dimensions[2].height = 34
merge_put(ws, 2, 2, 2, 4, "Daily Planner", font(TEXT, 24, True), PANEL, align=LEFT)
put(ws, 3, 2, "Date:", font(MUTED, 9), PANEL, align=LEFT)
input_cell(ws, 3, 3, None, DATE)
put(ws, 3, 4, None, fill_color=PANEL)
ws["N3"] = f'=IF(C3="",{TODAY},C3)'
ws["N3"].number_format = DATE
put(ws, 4, 2, '=TEXT($N$3,"dddd, mmmm d, yyyy")', font(TEAL, 11, True), BG, align=LEFT)
put(ws, 4, 3, None)
put(ws, 3, 7, "← leave blank for today", font(MUTED, 8, italic=True))
band(ws, 6, 2, 3, "TODAY'S SCHEDULE", LAVENDER)
for k in range(40):
    r = 7 + k
    put(ws, r, 2, f"={DAYSTART}+{k}*{INTERVAL}/1440", font(MUTED, 9), PANEL, "h:mm AM/PM", CENTER, BOTTOM)
    ws[f"D{r}"] = (f'=IFERROR(_xlfn.AGGREGATE(15,6,(ROW(Calc!$AW$2:$AW$2001)-1)/((Calc!$AW$2:$AW$2001=$N$3)*'
                   f'(Calc!$AX$2:$AX$2001<=B{r}+0.00001)*(Calc!$AY$2:$AY$2001>B{r}+0.00001)*'
                   f'(Calc!$AZ$2:$AZ$2001<>"Cancelled")),1),"")')
    put(ws, r, 3, f'=IF(D{r}="","",INDEX(Calc!$AU$2:$AU$2001,D{r}))', font(TEXT, 9, True), PANEL, align=LEFT,
        border=BOTTOM)
    for i in range(6):
        cf_formula(ws, f"C{r}", f'AND(ISNUMBER($D{r}),INDEX(TaskCat,$D{r})=INDEX({CATS},{i+1}))', CATCOLS[i])
ws.column_dimensions["D"].hidden = True
band(ws, 6, 6, 8, "TODAY'S TASKS", TEAL)
put(ws, 6, 9, f'=IFERROR(COUNTIFS(TaskDue,$N$3,TaskDone,1)/COUNTIFS(TaskDue,$N$3,TaskStat,"<>Cancelled"),0)',
    font(DARK, 12, True), TEAL, PCT, CENTER)
for i, h in enumerate(["", "TASK", "CATEGORY", "PRIORITY"]):
    put(ws, 7, 6 + i, h, font(MUTED, 8, True), PANEL, align=CENTER if i != 1 else LEFT)
task_rows(ws, 8, 25, lambda k: f'=IFERROR(MOD(_xlfn.AGGREGATE(15,6,Calc!$BN$2:$BN$2001/(Calc!$AW$2:$AW$2001=$N$3),{k}),10000),"")',
          [("F", "tick"), ("G", "AU"), ("H", "AV"), ("I", "BA")], "J")
for i in range(6):
    cf_formula(ws, "H8:H32", f'AND(H8<>"",H8=INDEX({CATS},{i+1}))', CATCOLS[i])
band(ws, 6, 12, 13, "MAIN FOCUS", PINK)
merge_put(ws, 7, 12, 9, 13, "Write project proposal and send before midday.", font(TEXT, 10), PANEL, align=WRAP)
band(ws, 11, 12, 13, "DAILY GRATITUDE", LAVENDER)
merge_put(ws, 12, 12, 14, 13, "=IFERROR(INDEX('Gratitude Log'!$C$6:$C$400,MATCH($N$3,'Gratitude Log'!$B$6:$B$400,0))&\"\",\"(add an entry in the Gratitude Log)\")",
          font(TEXT, 10, italic=True), PANEL, align=WRAP)
band(ws, 16, 12, 13, "DAILY HABITS", AMBER)
for i in range(HABN):
    r = 17 + i
    put(ws, r, 12, f"=IF('Habit Tracker'!$B${HAB0+i}=\"\",\"\",'Habit Tracker'!$B${HAB0+i})", font(TEXT, 9),
        PANEL, align=LEFT, border=BOTTOM)
    put(ws, r, 13, (f"=IF(L{r}=\"\",\"\",IFERROR(IF(INDEX('Habit Tracker'!$E${HAB0+i}:${col(HCOL0+NDAYS-1)}${HAB0+i},"
                    f"$N$3-'Habit Tracker'!$E$7+1)=\"✓\",\"✓\",\"○\"),\"\"))"), font(AMBER, 10, True), PANEL,
        align=CENTER, border=BOTTOM)
band(ws, 33, 12, 13, "NOTES", TEAL)
merge_put(ws, 34, 12, 46, 13, None, font(TEXT, 10), PANEL, align=WRAP)

# ============================================================================ Weekly
ws = sh["Weekly"]
wd = {"A": 2, "B": 10}
for d in range(7):
    wd[col(3 + 2 * d)] = 19
    wd[col(4 + 2 * d)] = 3
setup_sheet(ws, wd, tab=AMBER, zoom=85, rows=60)
ws.row_dimensions[2].height = 34
merge_put(ws, 2, 2, 2, 7, "Weekly Planner", font(TEXT, 24, True), PANEL, align=LEFT)
put(ws, 3, 2, "Start date:", font(MUTED, 9), PANEL, align=LEFT)
ws.merge_cells("C3:D3")
input_cell(ws, 3, 3, None, DATE)
put(ws, 3, 5, "← blank = current week (week start set on Setup)", font(MUTED, 8, italic=True))
ws["S3"] = f'=IF(C3="",{TODAY}-MOD(WEEKDAY({TODAY},IF({WEEKSTART}="Monday",2,1))-1,7),C3)'
put(ws, 4, 2, "Week of", font(MUTED, 9), BG)
put(ws, 4, 3, '=TEXT($S$3,"d mmm")&" – "&TEXT($S$3+6,"d mmm yyyy")', font(TEAL, 11, True), BG, align=LEFT)
dcol = [CORAL, AMBER, TEAL, BLUE, LAVENDER, PINK, GREEN]
for d in range(7):
    c = 3 + 2 * d
    h = col(c + 1)
    ws.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c + 1)
    put(ws, 6, c, f'=UPPER(TEXT($S$3+{d},"dddd"))', font(DARK, 10, True), dcol[d], align=CENTER)
    ws.merge_cells(start_row=7, start_column=c, end_row=7, end_column=c + 1)
    put(ws, 7, c, f'=TEXT($S$3+{d},"mmm d")', font(MUTED, 9, True), PANEL, align=CENTER)
    ws.merge_cells(start_row=8, start_column=c, end_row=8, end_column=c + 1)
    put(ws, 8, c, (f'=IFERROR(COUNTIFS(TaskDue,$S$3+{d},TaskDone,1)/COUNTIFS(TaskDue,$S$3+{d},TaskStat,"<>Cancelled"),0)'),
        font(TEXT, 12, True), PANEL, PCT, CENTER)
    databar(ws, f"{col(c)}8", dcol[d])
    for k in range(1, 13):
        r = 8 + k
        ws[f"{h}{r}"] = (f'=IFERROR(MOD(_xlfn.AGGREGATE(15,6,Calc!$BN$2:$BN$2001/(Calc!$AW$2:$AW$2001=$S$3+{d}),{k}),'
                         f'10000),"")')
        put(ws, r, c, f'=IF({h}{r}="","",INDEX(Calc!$AU$2:$AU$2001,{h}{r}))', font(TEXT, 8), PANEL, align=LEFT,
            border=BOTTOM)
        ws.cell(r, c + 1).border = BOTTOM
    cf_formula(ws, f"{col(c)}9:{col(c)}20", f'AND(ISNUMBER({h}9),INDEX(TaskDone,{h}9)=1)', "6E6E86")
    for i in range(6):
        cf_formula(ws, f"{col(c)}9:{col(c)}20", f'AND(ISNUMBER({h}9),INDEX(TaskCat,{h}9)=INDEX({CATS},{i+1}))',
                   CATCOLS[i])
band(ws, 22, 2, 16, "WEEKLY SCHEDULE  (auto-fills from task start / end times)", LAVENDER)
for d in range(7):
    c = 3 + 2 * d
    ws.merge_cells(start_row=23, start_column=c, end_row=23, end_column=c + 1)
    put(ws, 23, c, f'=TEXT($S$3+{d},"ddd d")', font(MUTED, 9, True), PANEL, align=CENTER)
for k in range(30):
    r = 24 + k
    put(ws, r, 2, f"={DAYSTART}+{k}*{INTERVAL}/1440", font(MUTED, 8), PANEL, "h:mm AM/PM", CENTER, BOTTOM)
    for d in range(7):
        c = 3 + 2 * d
        h = col(c + 1)
        ws[f"{h}{r}"] = (f'=IFERROR(_xlfn.AGGREGATE(15,6,(ROW(Calc!$AW$2:$AW$2001)-1)/((Calc!$AW$2:$AW$2001=$S$3+{d})*'
                         f'(Calc!$AX$2:$AX$2001<=$B{r}+0.00001)*(Calc!$AY$2:$AY$2001>$B{r}+0.00001)*'
                         f'(Calc!$AZ$2:$AZ$2001<>"Cancelled")),1),"")')
        put(ws, r, c, f'=IF({h}{r}="","",INDEX(Calc!$AU$2:$AU$2001,{h}{r}))', font(TEXT, 8), PANEL, align=LEFT,
            border=BOTTOM)
for d in range(7):
    c = 3 + 2 * d
    h = col(c + 1)
    for i in range(6):
        cf_formula(ws, f"{col(c)}24:{col(c)}53", f'AND(ISNUMBER({h}24),INDEX(TaskCat,{h}24)=INDEX({CATS},{i+1}))',
                   DARK, CATCOLS[i])
    ws.column_dimensions[h].hidden = True
ws.column_dimensions["S"].hidden = True

# ============================================================================ Monthly
ws = sh["Monthly"]
md = {"A": 2}
for d in range(7):
    md[col(2 + 2 * d)] = 20
    md[col(3 + 2 * d)] = 3
setup_sheet(ws, md, tab=PINK, zoom=85, rows=52)
merge_put(ws, 2, 2, 3, 6, '=TEXT($R$3,"mmmm yyyy")', font(TEXT, 24, True), BG, align=LEFT)
put(ws, 2, 8, "Any date in month:", font(MUTED, 9), PANEL, align=RIGHT)
ws.merge_cells("J2:K2")
input_cell(ws, 2, 10, None, DATE)
put(ws, 3, 8, "blank = this month", font(MUTED, 8, italic=True))
ws["R3"] = f'=DATE(YEAR(IF(J2="",{TODAY},J2)),MONTH(IF(J2="",{TODAY},J2)),1)'
ws["R4"] = f'=R3-MOD(WEEKDAY(R3,IF({WEEKSTART}="Monday",2,1))-1,7)'
for d in range(7):
    c = 2 + 2 * d
    ws.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c + 1)
    put(ws, 5, c, f'=UPPER(TEXT($R$4+{d},"dddd"))', font(DARK, 10, True), LAVENDER, align=CENTER)
for w in range(6):
    base = 6 + 7 * w
    for d in range(7):
        c = 2 + 2 * d
        h = col(c + 1)
        ws[f"{h}{base}"] = f"=$R$4+{7*w+d}"
        put(ws, base, c, f"=DAY({h}{base})", font(TEXT, 11, True), "2A2A3A", align=LEFT)
        cf_formula(ws, f"{col(c)}{base}", f"MONTH({h}{base})<>MONTH($R$3)", "5A5A70")
        cf_formula(ws, f"{col(c)}{base}", f"{h}{base}={TODAY}", DARK, CORAL)
        for k in range(1, 7):
            r = base + k
            ws[f"{h}{r}"] = (f'=IFERROR(MOD(_xlfn.AGGREGATE(15,6,Calc!$BN$2:$BN$2001/(Calc!$AW$2:$AW$2001={h}${base}),{k}),'
                             f'10000),"")')
            put(ws, r, c, f'=IF({h}{r}="","",INDEX(Calc!$AU$2:$AU$2001,{h}{r}))', font(TEXT, 8), PANEL,
                align=LEFT)
        ws.cell(base + 6, c).border = BOTTOM
for d in range(7):
    c = 2 + 2 * d
    h = col(c + 1)
    rng = f"{col(c)}7:{col(c)}47"
    cf_formula(ws, rng, f'AND(ISNUMBER({h}7),{h}7<5000,INDEX(TaskDone,MAX(1,MIN(2000,N({h}7))))=1)', "6E6E86", stop=True)
    for i in range(6):
        cf_formula(ws, rng, f'AND(ISNUMBER({h}7),{h}7<5000,INDEX(TaskCat,MAX(1,MIN(2000,N({h}7))))=INDEX({CATS},{i+1}))',
                   CATCOLS[i])
    ws.column_dimensions[h].hidden = True
ws.column_dimensions["R"].hidden = True
put(ws, 49, 2, "Up to 6 tasks per day, sorted by priority. Colour = category, grey = done.",
    font(MUTED, 8, italic=True))

# ============================================================================ Decision Matrix
ws = sh["Decision Matrix"]
setup_sheet(ws, {"A": 2, "B": 26, "C": 9, "D": 5, "E": 2, "F": 26, "G": 9, "H": 5, "I": 2, "J": 26, "K": 9,
                 "L": 5, "M": 2, "N": 26, "O": 9, "P": 5}, tab=CORAL, zoom=85, rows=45)
merge_put(ws, 2, 2, 2, 6, "Decision Matrix", font(TEXT, 22, True), BG, align=LEFT)
put(ws, 3, 2, "Eisenhower matrix: urgent = priority 1-Urgent / 2-High · important = ✓ in IMPORTANT.",
    font(MUTED, 9, italic=True))
flts = [("Start date", f"=DATE(YEAR({TODAY}),MONTH({TODAY}),1)", DATE), ("End date", f"=EOMONTH({TODAY},0)", DATE),
        ("Category", None, None), ("Person", None, None), ("Hide completed", "No", None)]
for i, (t, v, f) in enumerate(flts):
    put(ws, 2 + i, 10, t, font(MUTED, 9), PANEL, align=RIGHT)
    input_cell(ws, 2 + i, 11, v, f)
    ws.merge_cells(start_row=2 + i, start_column=11, end_row=2 + i, end_column=12)
dv_list(ws, f"={CATS}", "K4"); dv_list(ws, f"={PEOPLE}", "K5"); dv_list(ws, "=Lists!$K$2:$K$3", "K6")
cond = (f'(Calc!$AW$2:$AW$2001>=$K$2)*(Calc!$AW$2:$AW$2001<=$K$3)*(Calc!$BH$2:$BH$2001=1)*'
        f'(Calc!$AZ$2:$AZ$2001<>"Cancelled")*(($K$4="")+(Calc!$AV$2:$AV$2001=$K$4)>0)*'
        f'(($K$5="")+(Calc!$BC$2:$BC$2001=$K$5)>0)')
quads = [("DO FIRST", "Urgent & important — do it now.", CORAL), ("SCHEDULE", "Important, not urgent — plan it.", TEAL),
         ("DELEGATE", "Urgent, not important — hand it off.", AMBER), ("DELETE", "Neither — drop it.", LAVENDER)]
for q, (t, desc, colr) in enumerate(quads):
    c = 2 + 4 * q
    band(ws, 8, c, c + 2, t, colr, 12)
    put(ws, 9, c, desc, font(MUTED, 8, italic=True), PANEL, align=LEFT)
    ws.merge_cells(start_row=9, start_column=c, end_row=9, end_column=c + 2)
    put(ws, 10, c, "REMAINING", font(MUTED, 8, True), PANEL, align=LEFT)
    put(ws, 10, c + 1, "DONE", font(MUTED, 8, True), PANEL, align=CENTER)
    ws.cell(10, c + 2).fill = fill(PANEL)
    put(ws, 11, c, f'=SUMPRODUCT({cond}*(Calc!$BD$2:$BD$2001="{t}")*(Calc!$BG$2:$BG$2001=0))', font(colr, 20, True),
        PANEL, "0", LEFT)
    put(ws, 11, c + 1, f'=SUMPRODUCT({cond}*(Calc!$BD$2:$BD$2001="{t}")*(Calc!$BG$2:$BG$2001=1))',
        font(TEXT, 16, True), PANEL, "0", CENTER)
    ws.cell(11, c + 2).fill = fill(PANEL)
    ws.row_dimensions[11].height = 28
    ws.merge_cells(start_row=12, start_column=c, end_row=12, end_column=c + 2)
    put(ws, 12, c, f"=IFERROR({col(c+1)}11/({col(c)}11+{col(c+1)}11),0)", font(TEXT, 10, True), PANEL, PCT, CENTER)
    databar(ws, f"{col(c)}12", colr)
    put(ws, 13, c, "TASK", font(MUTED, 8, True), PANEL, align=LEFT)
    put(ws, 13, c + 1, "DAYS LEFT", font(MUTED, 8, True), PANEL, align=CENTER)
    put(ws, 13, c + 2, "", font(MUTED, 8, True), PANEL)
    h = col(c + 3) if q < 3 else "R"
    hc = f"{col(20 + q)}"
    for k in range(1, 21):
        r = 13 + k
        ws[f"{hc}{r}"] = (f'=IFERROR(MOD(_xlfn.AGGREGATE(15,6,((Calc!$AW$2:$AW$2001-40000)*10000+ROW(Calc!$AW$2:$AW$2001)-1)/'
                          f'({cond}*(Calc!$BD$2:$BD$2001="{t}")*(($K$6="No")+(Calc!$BG$2:$BG$2001=0)>0)),{k}),10000),"")')
        put(ws, r, c, f'=IF({hc}{r}="","",INDEX(Calc!$AU$2:$AU$2001,{hc}{r}))', font(TEXT, 9), PANEL, align=LEFT,
            border=BOTTOM)
        put(ws, r, c + 1, f'=IF({hc}{r}="","",INDEX(Calc!$AW$2:$AW$2001,{hc}{r})-{TODAY})', font(TEXT, 9), PANEL,
            "0", CENTER, BOTTOM)
        put(ws, r, c + 2, f'=IF({hc}{r}="","",IF(INDEX(Calc!$BG$2:$BG$2001,{hc}{r})=1,"✓",""))', font(colr, 10, True),
            PANEL, align=CENTER, border=BOTTOM)
    cf_formula(ws, f"{col(c)}14:{col(c+1)}33", f'AND(${hc}14<>"",INDEX(TaskDone,MAX(1,N(${hc}14)))=1)', "6E6E86")
    cf_formula(ws, f"{col(c+1)}14:{col(c+1)}33", f'AND(ISNUMBER({col(c+1)}14),{col(c+1)}14<0)', CORAL, bold=True)
for c in ("T", "U", "V", "W"):
    ws.column_dimensions[c].hidden = True
ws["Y10"] = "Quadrant"; ws["Z10"] = "Remaining"; ws["AA10"] = "Done"
for q, (t, _, _) in enumerate(quads):
    ws[f"Y{11+q}"] = t
    ws[f"Z{11+q}"] = f"={col(2+4*q)}11"
    ws[f"AA{11+q}"] = f"={col(3+4*q)}11"
band(ws, 35, 2, 8, "PROGRESS BY QUADRANT", LAVENDER)
bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100; bc.gapWidth = 50
bc.add_data(Reference(ws, min_col=27, min_row=10, max_row=14), titles_from_data=True)
bc.add_data(Reference(ws, min_col=26, min_row=10, max_row=14), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=25, min_row=11, max_row=14))
dark_chart(bc, legend_pos="t"); color_series(bc.series[0], TEAL); color_series(bc.series[1], "5C5C78")
bc.width, bc.height = 17, 6.5
ws.add_chart(bc, "B36")

# ============================================================================ Habit Tracker
ws = sh["Habit Tracker"]
hw = {"A": 2, "B": 30, "C": 8, "D": 7}
for d in range(NDAYS):
    hw[col(HCOL0 + d)] = 3.3
setup_sheet(ws, hw, tab=AMBER, zoom=85, extra_cols=3, rows=40)
merge_put(ws, 2, 2, 2, 4, "Habit Tracker", font(TEXT, 22, True), BG, align=LEFT)
put(ws, 3, 2, "Enter habits in column B, tick ✓ each day. 52 weeks in one tab — hide finished weeks if you like.",
    font(MUTED, 8, italic=True))
put(ws, 4, 2, "YEAR PROGRESS", font(MUTED, 8, True), PANEL, align=LEFT)
lastc = col(HCOL0 + NDAYS - 1)
put(ws, 4, 3, (f'=IFERROR(COUNTIF(E{HAB0}:{lastc}{HAB0+HABN-1},"✓")/(COUNTA(B{HAB0}:B{HAB0+HABN-1})*'
               f'MAX(1,MIN({NDAYS},{TODAY}-$E$7+1))),0)'), font(AMBER, 14, True), PANEL, PCT, CENTER)
ws["E7"] = f"=Setup!$D$5-MOD(WEEKDAY(Setup!$D$5,2)-1,7)"
wcols = [LAVENDER, TEAL, AMBER, PINK, BLUE, CORAL, GREEN, ROSE]
put(ws, 5, 2, "Month", font(MUTED, 8), PANEL); put(ws, 6, 2, "Week", font(MUTED, 8), PANEL)
put(ws, 7, 2, "Date", font(MUTED, 8), PANEL); put(ws, 8, 2, "HABITS", font(MUTED, 8, True), PANEL)
put(ws, 8, 3, "% DONE", font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 8, 4, "TOTAL", font(MUTED, 8, True), PANEL, align=CENTER)
for w in range(52):
    c0 = HCOL0 + 7 * w
    colr = wcols[w % 8]
    ws.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + 6)
    put(ws, 5, c0, f'=TEXT($E$7+{7*w},"mmmm")', font(MUTED, 8, True), PANEL, align=CENTER)
    band(ws, 6, c0, c0 + 6, f"WEEK {w+1}", colr, 8)
    for d in range(7):
        c = c0 + d
        dc = col(c)
        if w or d:
            ws.cell(7, c).value = f"=$E$7+{7*w+d}"
        ws.cell(7, c).number_format = "d"
        ws.cell(7, c).font = font(TEXT, 7, True)
        ws.cell(7, c).alignment = CENTER
        ws.cell(7, c).fill = fill(PANEL)
        put(ws, 8, c, f'=LEFT(TEXT({dc}7,"ddd"),1)', font(MUTED, 7), PANEL, align=CENTER)
        for i in range(HABN):
            cell = ws.cell(HAB0 + i, c)
            cell.fill = fill(ROW if i % 2 else ROW2)
            cell.alignment = CENTER
            cell.font = font(DARK, 8, True)
            cell.border = Border(left=Side(style="thin", color="2A2A38"))
        put(ws, 24, c, (f'=IF({dc}7>{TODAY},"",IFERROR(COUNTIF({dc}{HAB0}:{dc}{HAB0+HABN-1},"✓")/'
                        f'COUNTA($B${HAB0}:$B${HAB0+HABN-1}),0))'), font(MUTED, 6), PANEL, '0%', CENTER)
        put(ws, 28, c, None, font(DARK, 8, True), ROW, align=CENTER)
    ws.merge_cells(start_row=25, start_column=c0, end_row=25, end_column=c0 + 6)
    c1 = col(c0 + 6)
    put(ws, 25, c0, (f'=IF({col(c0)}7>{TODAY},"",IFERROR(COUNTIF({col(c0)}{HAB0}:{c1}{HAB0+HABN-1},"✓")/'
                     f'(7*COUNTA($B${HAB0}:$B${HAB0+HABN-1})),0))'), font(colr, 9, True), PANEL, PCT, CENTER)
    ws.merge_cells(start_row=29, start_column=c0, end_row=29, end_column=c0 + 6)
    put(ws, 29, c0, f'=IFERROR(AVERAGE({col(c0)}28:{c1}28),"")', font(colr, 9, True), PANEL, '0.0', CENTER)
    for i in range(5):
        r = 33 + i
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 6)
        put(ws, r, c0, None, font(DARK, 10, True), ROW if i % 2 else ROW2, align=CENTER)
        cf_formula(ws, f"{col(c0)}{r}", f'{col(c0)}{r}="✓"', DARK, colr)
    rng = f"{col(c0)}{HAB0}:{c1}{HAB0+HABN-1}"
    cf_formula(ws, rng, f'{col(c0)}{HAB0}="✓"', DARK, colr)
    # helper row for chart (week label / %)
    ws.cell(40, 2 + w).value = f"=\"W{w+1}\""
    ws.cell(41, 2 + w).value = f"=N({col(c0)}25)"
    ws.cell(42, 2 + w).value = f"=N({col(c0)}29)"
habits = ["Sleep 7 hours", "Make up bed", "10 min Meditation", "Drink 2L of water", "15 min stretching",
          "Walk min 7000 steps", "50 x Push Ups", "Eat 5 serves of fruit & veg", "No coffee after 2pm",
          "No eating after 7pm", "Read for 20 min", "Gratitude list: write 3 things", "Positive affirmations",
          "Don't use phone before bed", "In bed before 10pm"]
random.seed(11)
first_day = START_DATE - dt.timedelta(days=START_DATE.weekday())
today = D(2026, 9, 29)
for i in range(HABN):
    r = HAB0 + i
    input_cell(ws, r, 2, habits[i], align=LEFT)
    put(ws, r, 3, (f'=IF(B{r}="","",IFERROR(COUNTIF(E{r}:{lastc}{r},"✓")/MAX(1,MIN({NDAYS},{TODAY}-$E$7+1)),0))'),
        font(TEXT, 9, True), PANEL, PCT, CENTER)
    put(ws, r, 4, f'=IF(B{r}="","",COUNTIF(E{r}:{lastc}{r},"✓"))', font(MUTED, 9), PANEL, "0", CENTER)
    databar(ws, f"C{r}", AMBER)
    ndone = (today - first_day).days
    for d in range(ndone):
        if random.random() < 0.55 + 0.03 * (i % 5):
            ws.cell(r, HCOL0 + d).value = "✓"
for d in range((today - first_day).days):
    ws.cell(28, HCOL0 + d).value = random.choice([3, 4, 4, 5, 3, 2, 4, 5])
dv_list(ws, "=Lists!$D$2", f"E{HAB0}:{lastc}{HAB0+HABN-1}")
put(ws, 24, 2, "DAILY PROGRESS", font(MUTED, 8, True), PANEL, align=LEFT)
put(ws, 25, 2, "WEEKLY PROGRESS", font(MUTED, 8, True), PANEL, align=LEFT)
band(ws, 27, 2, 4, "MOOD TRACKER", PINK, 9)
put(ws, 28, 2, "Mood score (1 = really bad … 5 = really good)", font(TEXT, 8), PANEL, align=LEFT)
put(ws, 28, 3, f"=IFERROR(AVERAGE(E28:{lastc}28),\"\")", font(PINK, 12, True), PANEL, '0.0', CENTER)
put(ws, 28, 4, "yearly avg", font(MUTED, 7), PANEL, align=CENTER)
put(ws, 29, 2, "Weekly average", font(MUTED, 8), PANEL, align=LEFT)
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
ws.add_data_validation(dv); dv.add(f"E28:{lastc}28")
from openpyxl.formatting.rule import ColorScaleRule
ws.conditional_formatting.add(f"E28:{lastc}28", ColorScaleRule(start_type="num", start_value=1, start_color=CORAL,
                                                               mid_type="num", mid_value=3, mid_color=AMBER,
                                                               end_type="num", end_value=5, end_color=TEAL))
band(ws, 32, 2, 4, "WEEKLY HABITS  (tick ✓ once per week)", TEAL, 9)
wh = ["Weekly meal planning", "Review to-do list", "Review budget", "Call family", "Deep clean one room"]
for i in range(5):
    input_cell(ws, 33 + i, 2, wh[i], align=LEFT)
    put(ws, 33 + i, 3, f'=IF(B{33+i}="","",COUNTIF(E{33+i}:{lastc}{33+i},"✓"))', font(TEXT, 9, True), PANEL, "0",
        CENTER)
    for w in range(6):
        if random.random() < 0.7:
            ws.cell(33 + i, HCOL0 + 7 * w).value = "✓"
dv_list(ws, "=Lists!$D$2", f"E33:{lastc}37")
for r in (40, 41, 42):
    ws.row_dimensions[r].hidden = True
cf_formula(ws, f"E7:{lastc}8", f"E$7={TODAY}", DARK, CORAL)
ws.freeze_panes = "E9"
ch = BarChart(); ch.type = "col"; ch.gapWidth = 30
ch.add_data(Reference(ws, min_col=2, max_col=53, min_row=41), from_rows=True)
ch.set_categories(Reference(ws, min_col=2, max_col=53, min_row=40))
dark_chart(ch, legend=False, gridlines=True); color_series(ch.series[0], AMBER)
ch.y_axis.numFmt = "0%"; ch.y_axis.scaling.max = 1; ch.y_axis.scaling.min = 0
ch.visible_cells_only = False
ch.width, ch.height = 22, 6
band(ws, 44, 2, 4, "WEEKLY HABIT PROGRESS", AMBER, 9)
ws.add_chart(ch, "B46")
mc = BarChart(); mc.type = "col"; mc.gapWidth = 30
mc.add_data(Reference(ws, min_col=2, max_col=53, min_row=42), from_rows=True)
mc.set_categories(Reference(ws, min_col=2, max_col=53, min_row=40))
dark_chart(mc, legend=False, gridlines=True); color_series(mc.series[0], PINK)
mc.y_axis.scaling.max = 5; mc.y_axis.scaling.min = 0
mc.visible_cells_only = False
mc.width, mc.height = 22, 6
band(ws, 58, 2, 4, "WEEKLY MOOD", PINK, 9)
ws.add_chart(mc, "B60")

# ============================================================================ Gratitude Log
ws = sh["Gratitude Log"]
setup_sheet(ws, {"A": 2, "B": 13, "C": 90}, tab=LAVENDER, rows=10)
merge_put(ws, 2, 2, 2, 3, "Gratitude Log", font(TEXT, 20, True), BG, align=LEFT)
put(ws, 3, 2, "One line a day — today's entry shows up on the Daily planner.", font(MUTED, 9, italic=True))
band(ws, 4, 2, 3, "TODAY I'M GRATEFUL FOR…", LAVENDER)
put(ws, 5, 2, "DATE", font(MUTED, 8, True), PANEL, align=CENTER)
put(ws, 5, 3, "ENTRY", font(MUTED, 8, True), PANEL, align=LEFT)
grat = ["A slow coffee and a clear head this morning.", "Lunch with Amy and a proper laugh.",
        "Finishing the proposal draft early.", "Sam helping with dinner without being asked.",
        "A sunny walk at lunchtime.", "Clear direction and a productive mindset.", "My health and a good night's sleep."]
for i in range(395):
    r = 6 + i
    v = (D(2026, 9, 23) + dt.timedelta(days=i), grat[i]) if i < len(grat) else (None, None)
    input_cell(ws, r, 2, v[0], DATE)
    input_cell(ws, r, 3, v[1], align=LEFT)
cf_formula(ws, "B6:C400", f"$B6={TODAY}", DARK, TEAL)
ws.freeze_panes = "B6"

# ============================================================================ Goals
ws = sh["Goals"]
setup_sheet(ws, {"A": 2, "B": 16, "C": 16, "D": 16, "E": 12, "F": 2, "G": 30, "H": 12, "I": 11, "J": 7, "K": 12},
            tab=GREEN, zoom=85, rows=150)
merge_put(ws, 2, 2, 2, 6, "Goal Setting", font(TEXT, 22, True), BG, align=LEFT)
put(ws, 3, 2, '"Goals are dreams with deadlines." — Diana Scharf', font(MUTED, 9, italic=True))
band(ws, 5, 2, 5, "OVERALL PROGRESS", GREEN)
merge_put(ws, 6, 2, 8, 5, "=IFERROR(AVERAGEIF(K22:K250,\">=0\"),0)", font(GREEN, 28, True), PANEL, PCT)
band(ws, 5, 7, 11, "GOALS AT A GLANCE", LAVENDER)
gcat = [CORAL, AMBER, TEAL, PINK, LAVENDER, BLUE, GREEN, ROSE]
samples = [("Personal Growth", "Learn to speak Spanish", "To connect with more people and cultures",
            "Time management", "Trip to Spain", ["Find a tutor", "Use a language app daily", "Join a conversation group",
                                                 "Watch shows in Spanish", "Practise with a partner", "Book a trip"]),
           ("Career", "Promotion to Project Lead", "To advance my career and take on more responsibility",
            "Confidence presenting", "Weekend away", ["Talk to my manager", "Take a leadership course",
                                                     "Lead one project end-to-end", "Mentor a junior", "", ""]),
           ("Financial", "Save a 20% house deposit", "To own a home and build equity", "Impulse spending",
            "Housewarming party", ["Set up automatic transfers", "Cut 2 subscriptions", "Side-hustle income",
                                   "Review budget monthly", "", ""])]
NG = 12
for g in range(NG):
    r = 22 + g * 9
    colr = gcat[g % len(gcat)]
    s = samples[g] if g < len(samples) else (None, None, None, None, None, [""] * 6)
    band(ws, r, 2, 11, f"GOAL {g+1}", colr, 10)
    put(ws, r, 2, f'="GOAL {g+1}"&IF(C{r+1}="","","  ·  "&UPPER(C{r+1}))', font(DARK, 10, True), colr, align=LEFT)
    for i, lab in enumerate(["Category", "Goal", "Why", "Biggest obstacle", "Reward", "Deadline", "Notes"]):
        rr = r + 1 + i
        put(ws, rr, 2, lab.upper(), font(MUTED, 8, True), PANEL, align=LEFT)
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=5)
        v = [s[0], s[1], s[2], s[3], s[4], D(2027, 6, 30) if s[0] else None, None][i]
        input_cell(ws, rr, 3, v, DATE if lab == "Deadline" else None, LEFT)
        for c in (4, 5):
            ws.cell(rr, c).fill = fill(INPUT)
    put(ws, r + 1, 7, "STEP BY STEP TO GET THERE", font(MUTED, 8, True), PANEL, align=LEFT)
    ws.merge_cells(start_row=r + 1, start_column=7, end_row=r + 1, end_column=8)
    put(ws, r + 1, 9, "DUE", font(MUTED, 8, True), PANEL, align=CENTER)
    put(ws, r + 1, 10, "DONE", font(MUTED, 8, True), PANEL, align=CENTER)
    put(ws, r + 1, 11, "PROGRESS", font(MUTED, 8, True), PANEL, align=CENTER)
    for i in range(6):
        rr = r + 2 + i
        ws.merge_cells(start_row=rr, start_column=7, end_row=rr, end_column=8)
        step = s[5][i] if s[5] else None
        input_cell(ws, rr, 7, step or None, align=LEFT)
        ws.cell(rr, 8).fill = fill(INPUT)
        input_cell(ws, rr, 9, D(2026, 10 + i // 2, 15) if step else None, DATE)
        input_cell(ws, rr, 10, "✓" if step and i < (3 - g) else None)
    dv_list(ws, "=Lists!$D$2", f"J{r+2}:J{r+7}")
    merge_put(ws, r + 2, 11, r + 7, 11,
              f'=IF(COUNTA(G{r+2}:G{r+7})=0,"",COUNTIF(J{r+2}:J{r+7},"✓")/COUNTA(G{r+2}:G{r+7}))',
              font(colr, 20, True), PANEL, PCT)
    # summary row
    sr = 6 + g // 1 if False else None
for g in range(NG):
    r = 22 + g * 9
    rr = 6 + (g % 6)
    cc = 7 if g < 6 else 9
ws["M5"] = "Goal"; ws["N5"] = "Progress"
for g in range(NG):
    r = 22 + g * 9
    ws[f"M{6+g}"] = f'=IF(C{r+2}="","Goal {g+1}",C{r+2})'
    ws[f"N{6+g}"] = f"=N(K{r+2})"
ws.column_dimensions["M"].hidden = True
ws.column_dimensions["N"].hidden = True
gc = BarChart(); gc.type = "col"; gc.gapWidth = 50
gc.add_data(Reference(ws, min_col=14, min_row=5, max_row=5 + NG), titles_from_data=True)
gc.set_categories(Reference(ws, min_col=13, min_row=6, max_row=5 + NG))
dark_chart(gc, legend=False, gridlines=True); color_points(gc.series[0], gcat * 2)
gc.y_axis.numFmt = "0%"; gc.y_axis.scaling.max = 1; gc.y_axis.scaling.min = 0
gc.visible_cells_only = False
gc.width, gc.height = 15, 7.6
ws.add_chart(gc, "G6")
put(ws, 10, 2, "Tip: paste an inspiring image next to each goal (Insert › Pictures).", font(MUTED, 8, italic=True))

# ============================================================================ Instructions
ws = sh["Instructions"]
setup_sheet(ws, {"A": 3, "B": 4, "C": 110}, tab=LAVENDER, rows=50)
merge_put(ws, 2, 2, 2, 3, "TASK + HABIT TRACKER", font(TEXT, 24, True), BG, align=LEFT)
put(ws, 3, 3, "Tasks · planners · decision matrix · habits · mood · gratitude · goals  —  dark edition",
    font(TEAL, 11, True))
steps = [
    ("SETUP", "Set the planner start date (habit tracker & 1-year recurring horizon), schedule interval, day start "
              "time, week start and your categories & people. Yellow cells are inputs everywhere."),
    ("RECURRING SCHEDULE", "Set up repeating tasks ONCE (daily, weekdays, weekly, bi-weekly, monthly, quarterly, annual "
                           "or once) with times, priority, importance and person. They auto-fill the Recurring Log."),
    ("VARIABLE TASKS", "Type one-off tasks here. DAYS LEFT, DURATION and DECISION (Eisenhower) calculate automatically."),
    ("RECURRING LOG", "Every occurrence of your recurring tasks. Update STATUS, PROGRESS and NOTES only. If you change "
                      "the schedule later, re-check statuses (rows re-sort by date)."),
    ("TASKS FILTER", "All tasks (recurring + variable) in one view. Sort by due date / priority / category / person / "
                     "status and filter by date range, name, priority, importance, category, status, person and "
                     "decision. Click ✎ edit to jump to the exact row."),
    ("DASHBOARD", "Today's tasks, overdue tasks, this week's tasks and a summary (decisions, status, category, person) "
                  "with its own date / category / person filters. Links on the left jump to every tab."),
    ("DAILY / WEEKLY / MONTHLY", "Leave the date blank for today / this week / this month or pick any date. Tasks "
                                 "auto-sort by priority, the schedule fills from start & end times, colours follow "
                                 "category and done tasks turn grey."),
    ("DECISION MATRIX", "Do first / schedule / delegate / delete with remaining vs done counts, progress bars and task "
                        "lists. Urgent = priority 1-Urgent or 2-High; important = ✓."),
    ("HABIT TRACKER", "15 daily habits over 52 weeks, 5 weekly habits and a 1–5 mood tracker, with weekly progress "
                      "and mood charts. Hide finished weeks' columns to keep it tidy."),
    ("GRATITUDE & GOALS", "Daily gratitude log (shows on the Daily planner) and 12 goals with why / obstacle / reward "
                          "and 6 steps each."),
    ("NOTES", "Built for Microsoft Excel 365 / 2021+. Sample data is included — overwrite or delete it."),
]
r = 5
for t, txt in steps:
    put(ws, r, 3, t, font(AMBER, 11, True))
    put(ws, r + 1, 3, txt, font(TEXT, 10), align=WRAP)
    ws.row_dimensions[r + 1].height = 30
    r += 3

wb.active = 0
wb.calculation.fullCalcOnLoad = True
out = "Task_Habit_Tracker.xlsx"
wb.save(out)
print("saved", out)
