"""Shared styling helpers for the dark-theme planner workbooks."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"

# Palette
BG = "121218"        # page background
PANEL = "1C1C27"     # card background
ROW = "23232F"       # table row
ROW2 = "1C1C27"
LINE = "3A3A4C"
TEXT = "ECECF4"
MUTED = "9C9CB4"
DARK = "121218"

TEAL = "35D6BE"      # income
AMBER = "FFB224"     # savings
PINK = "FF7AB8"      # bills
CORAL = "FF4F8B"     # debt
ROSE = "F6A6D1"      # subscriptions
LAVENDER = "9D8CF7"  # variable
BLUE = "5AB4FF"
GREEN = "7BE08A"
INPUT = "FFF3B0"     # input cells (light yellow on dark theme)

CAT_COLORS = {
    "Income": TEAL, "Savings": AMBER, "Bills": PINK, "Debt": CORAL,
    "Subscriptions": ROSE, "Variable Exp": LAVENDER, "Transfer": BLUE,
}

MONEY = '$#,##0.00;-$#,##0.00;"-"'
MONEY0 = '$#,##0;-$#,##0;"-"'
PCT = '0%;-0%;"0%"'
DATE = 'd-mmm-yy'

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def font(color=TEXT, size=10, bold=False, italic=False, name=FONT):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
RIGHT = Alignment(horizontal="right", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)


def setup_sheet(ws, widths, tab=None, zoom=90, extra_cols=10, rows=None):
    """Dark background everywhere, no gridlines, set column widths."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    if tab:
        ws.sheet_properties.tabColor = tab
    last = 0
    for col, w in widths.items():
        idx = col if isinstance(col, int) else None
        if idx is None:
            from openpyxl.utils import column_index_from_string
            idx = column_index_from_string(col)
        ws.column_dimensions[get_column_letter(idx)].width = w
        last = max(last, idx)
    for i in range(1, last + extra_cols + 1):
        cd = ws.column_dimensions[get_column_letter(i)]
        cd.fill = fill(BG)
        cd.font = font()
    if rows:
        paint(ws, 1, 1, rows, last + extra_cols, BG)


def paint(ws, r1, c1, r2, c2, color, fnt=None):
    f = fill(color)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.fill = f
            if fnt is not None:
                cell.font = fnt
            elif cell.font is None or cell.font.name != FONT:
                cell.font = font()


def put(ws, r, c, value, fnt=None, fill_color=None, fmt=None, align=None, border=None):
    cell = ws.cell(r, c)
    cell.value = value
    cell.font = fnt or font()
    if fill_color:
        cell.fill = fill(fill_color)
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell


def band(ws, r, c1, c2, text, color, size=11, text_color=DARK, height=None):
    """Coloured header band spanning c1..c2."""
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    paint(ws, r, c1, r, c2, color)
    put(ws, r, c1, text, font(text_color, size, bold=True), color, align=CENTER)
    if height:
        ws.row_dimensions[r].height = height


def merge_put(ws, r1, c1, r2, c2, value, fnt=None, fill_color=None, fmt=None, align=CENTER):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    if fill_color:
        paint(ws, r1, c1, r2, c2, fill_color)
    return put(ws, r1, c1, value, fnt, fill_color, fmt, align)


def input_cell(ws, r, c, value=None, fmt=None, align=CENTER):
    cell = put(ws, r, c, value, font("121218", 10, bold=True), INPUT, fmt, align, BORDER)
    return cell


def dv_list(ws, formula, rng, allow_blank=True, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
    dv.error = "Please pick a value from the list."
    dv.errorStyle = "warning"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def cf_formula(ws, rng, formula, font_color=None, fill_color=None, bold=None, stop=False):
    kw = {}
    if font_color:
        kw["font"] = Font(color=font_color, bold=bold)
    if fill_color:
        kw["fill"] = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        kw["fill"].bgColor = fill_color
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], stopIfTrue=stop, **kw))


def databar(ws, rng, color, max_value=1):
    ws.conditional_formatting.add(
        rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max_value,
                         color=color, showValue=True))


# ---------------------------------------------------------------- charts
def _txt(color=TEXT, size=900, bold=False):
    cp = CharacterProperties(solidFill=color, sz=size, b=bold, latin=None)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def dark_chart(chart, legend=True, legend_pos="r", axes=True, gridlines=False):
    chart.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    chart.roundedCorners = False
    if legend and chart.legend is not None:
        chart.legend.position = legend_pos
        chart.legend.txPr = _txt()
    else:
        chart.legend = None
    if axes:
        for ax in (chart.x_axis, chart.y_axis):
            ax.delete = False
            ax.txPr = _txt(MUTED, 800)
            ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
        if gridlines:
            chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="2C2C3A"))
        else:
            chart.y_axis.majorGridlines = None
    return chart


def color_series(series, color, line=False):
    if line:
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 28000
        series.smooth = False
    else:
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.noFill = True


def color_points(series, colors):
    pts = []
    for i, c in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = c
        pt.graphicalProperties.line.solidFill = PANEL
        pts.append(pt)
    series.dPt = pts


def pct_labels(chart):
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLeaderLines = False
    chart.dataLabels.txPr = _txt(DARK, 800, True)


def value_labels(series, color=TEXT, fmt=None):
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    for a in ("showPercent", "showCatName", "showSerName", "showLegendKey"):
        setattr(series.dLbls, a, False)
    if fmt:
        series.dLbls.numFmt = fmt
    series.dLbls.txPr = _txt(color, 700)


PALETTE = [TEAL, PINK, AMBER, LAVENDER, BLUE, CORAL, GREEN, ROSE, "FF9F5A", "C084FC",
           "4ADE80", "F472B6", "38BDF8", "FACC15", "A3A3FF"]


def col(c):
    return get_column_letter(c)
